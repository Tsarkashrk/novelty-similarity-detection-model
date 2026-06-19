from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "Hybrid_Thesis_Novelty_Article.docx"


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x55, 0x55, 0x55)
FILL = "F4F6F9"
HEADER_FILL = "F4F6F9"
BORDER = "B8C2CC"


def set_font(run, name: str = "Calibri", size: float | None = None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Calibri", size=11, color=INK, bold=False):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:cs"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent_dxa: int = 120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.append(tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = borders.find(qn(f"w:{name}"))
        if border is None:
            border = OxmlElement(f"w:{name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)


def format_table_text(table, header_rows=1):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    set_font(run, size=10, color=INK, bold=(r_idx < header_rows))
            if r_idx < header_rows:
                set_cell_shading(cell, HEADER_FILL)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=10, color=MUTED, italic=True)
    return p


def add_para(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.333
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_font(r)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_font(r)
    return p


def add_callout(doc: Document, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=10.5, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_font(r2, size=10)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[int]):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].text = value
    set_table_geometry(table, widths)
    set_table_borders(table)
    format_table_text(table)
    doc.add_paragraph()
    return table


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style_font(styles["Normal"], size=11)
    styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.line_spacing = 1.333

    set_style_font(styles["Title"], size=18, color=INK, bold=True)
    styles["Title"].paragraph_format.space_after = Pt(6)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    set_style_font(styles["Subtitle"], size=11, color=MUTED)
    styles["Subtitle"].paragraph_format.space_after = Pt(12)
    styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        set_style_font(style, size=11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run("Hybrid Thesis Topic Novelty Assessment")
    set_font(r, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Draft manuscript for academic revision")
    set_font(r, size=9, color=MUTED)


def build_article():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(
        "Hybrid Semantic Novelty Assessment for Thesis Topic Management: "
        "A Reproducible Model for Duplicate Detection in Academic Project Platforms"
    )
    set_font(r, size=18, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("[Student Name], [University], [Department] | Original research article draft")
    set_font(r, size=11, color=MUTED)

    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "Universities increasingly manage diploma and master thesis projects through digital platforms, "
        "yet topic approval is often still based on manual checking, informal supervisor memory, or exact "
        "keyword matching. This creates a risk of duplicated topics, weakly differentiated research plans, "
        "and uneven supervisor allocation. This study proposes a hybrid semantic model for thesis topic "
        "novelty assessment and duplicate detection. The model combines multilingual sentence embeddings, "
        "cosine similarity search, keyword uniqueness, domain rarity, novelty scoring, and threshold-based "
        "duplicate classification. The proposed pipeline is designed for multilingual academic environments "
        "where Russian, English, Kazakh, and mixed-language thesis titles may occur in the same repository. "
        "The novelty score is formulated as a weighted combination of semantic distance, inverse keyword "
        "frequency, and domain-specific rarity. The architecture is compatible with Sentence-BERT, "
        "multilingual E5 embeddings, exact cosine search for small repositories, and PostgreSQL with pgvector "
        "for scalable deployment. A small prototype validation confirms computational consistency of the "
        "classification workflow; however, a full institutional evaluation with expert-labelled thesis-topic "
        "pairs is required before submission. The main contribution is a reproducible and explainable model "
        "that transforms topic approval from subjective manual comparison into measurable decision support "
        "for academic project management platforms.",
    )
    add_para(
        doc,
        "Keywords: thesis topic management; novelty assessment; duplicate detection; semantic similarity; "
        "Sentence-BERT; multilingual E5; academic project management; pgvector.",
    )

    add_heading(doc, "1. Introduction", 1)
    add_para(
        doc,
        "The management of diploma and master thesis projects is a recurring administrative and academic "
        "challenge in higher education. Each academic year, departments approve large numbers of student "
        "research topics, assign supervisors, monitor milestones, and evaluate whether proposed work is "
        "sufficiently original. In practice, topic approval often depends on manual review by programme "
        "coordinators or supervisors. This process is slow, difficult to reproduce, and vulnerable to "
        "inconsistency when topic repositories contain hundreds or thousands of historical records.",
    )
    add_para(
        doc,
        "A particularly important problem is the detection of duplicated or near-duplicated thesis topics. "
        "Exact lexical matching is insufficient because two topics can express the same research intention "
        "with different wording, different languages, or different levels of specificity. For example, "
        "'semantic similarity model for thesis duplicate detection' and 'intelligent system for identifying "
        "similar diploma topics' may be semantically close even if their surface forms differ. Conversely, "
        "two titles can share frequent academic words such as system, platform, management, or artificial "
        "intelligence while addressing substantially different research problems.",
    )
    add_para(
        doc,
        "Recent progress in transformer-based language models has made semantic text representation more "
        "practical for such tasks. BERT introduced deep bidirectional contextual representations [2], while "
        "Sentence-BERT adapted transformer encoders for efficient sentence-level similarity comparison [3]. "
        "The E5 family further demonstrated strong general-purpose text embeddings for retrieval, clustering, "
        "and classification tasks [4]. These developments make it possible to represent thesis topics as dense "
        "vectors and compare them through cosine similarity. However, semantic similarity alone does not fully "
        "capture novelty: a topic may be semantically close to previous work but still novel because it uses a "
        "rare method, dataset, language, or domain combination. Therefore, novelty assessment requires a hybrid "
        "model rather than a single similarity score.",
    )
    add_para(
        doc,
        "This paper addresses the following research question: how can a reproducible and explainable model "
        "assess the novelty of thesis topics and detect duplicates in a multilingual academic project "
        "management platform? The objective is to design a model that is technically feasible for a master "
        "research project, sufficiently modern for publication-oriented work, and transparent enough for "
        "academic administrators and supervisors to interpret.",
    )
    add_para(doc, "The contributions of this study are as follows:")
    add_number(
        doc,
        "A hybrid thesis topic novelty assessment model combining semantic embeddings, keyword uniqueness, "
        "domain rarity, and duplicate classification.",
    )
    add_number(
        doc,
        "A reproducible scoring formula that distinguishes semantic similarity from research novelty.",
    )
    add_number(
        doc,
        "A deployment-oriented architecture compatible with FastAPI, PostgreSQL, pgvector, and exact vector "
        "search for small to medium university repositories.",
    )
    add_number(
        doc,
        "An evaluation protocol including baselines, threshold optimisation, ablation study, and expert-labelled "
        "validation pairs.",
    )

    add_heading(doc, "2. Related Work", 1)
    add_heading(doc, "2.1 Semantic Text Representation", 2)
    add_para(
        doc,
        "Transformer architectures have become a foundation for modern natural language processing. The "
        "original Transformer introduced attention-based sequence modelling without recurrence [1], and BERT "
        "extended this direction through bidirectional pre-training [2]. Although BERT is effective for many "
        "language understanding tasks, directly comparing raw BERT sentence representations is computationally "
        "inefficient for large pairwise similarity search. Sentence-BERT addresses this limitation by using "
        "siamese and triplet network structures to produce sentence embeddings that can be compared efficiently "
        "with cosine similarity [3]. This makes SBERT-style models suitable for semantic duplicate detection, "
        "retrieval, and clustering.",
    )
    add_para(
        doc,
        "For multilingual academic environments, multilingual embedding models are preferable to English-only "
        "models. The E5 family provides general-purpose embeddings trained through weakly supervised contrastive "
        "pre-training and has been reported as effective across retrieval and classification tasks [4]. "
        "Sentence-transformer models such as paraphrase-multilingual-MiniLM-L12-v2 offer a lighter alternative "
        "for multilingual sentence similarity where computational resources are limited [9].",
    )
    add_heading(doc, "2.2 Semantic Search and Vector Storage", 2)
    add_para(
        doc,
        "Once text is converted into embeddings, nearest-neighbour search can be implemented through exact "
        "cosine similarity for small repositories or approximate indexing for larger collections. FAISS "
        "demonstrated efficient billion-scale similarity search using GPU-oriented algorithms [5]. In academic "
        "platform prototypes, however, a simpler architecture may be sufficient when the dataset contains a few "
        "hundred or a few thousand thesis topics. PostgreSQL with pgvector is a practical deployment option "
        "because embeddings, metadata, supervisor records, and similarity decisions can be stored in the same "
        "relational database while still supporting vector similarity search [8].",
    )
    add_heading(doc, "2.3 Lexical Signals and Novelty", 2)
    add_para(
        doc,
        "Dense semantic representations are powerful, but they can hide interpretable lexical signals. "
        "Information retrieval research has long used term weighting and probabilistic relevance models such "
        "as BM25 to estimate the importance of query terms in a document collection [6]. In the context of "
        "thesis topic evaluation, rare keywords can indicate a potentially novel method, dataset, or domain, "
        "whereas highly frequent terms such as platform, management, or artificial intelligence contribute "
        "little evidence of originality. Therefore, keyword uniqueness should complement semantic similarity.",
    )
    add_heading(doc, "2.4 Research Gap", 2)
    add_para(
        doc,
        "Existing approaches to academic project administration often focus on workflow management, document "
        "submission, or supervisor allocation. At the same time, semantic duplicate detection methods often "
        "focus on document-level plagiarism or general paraphrase similarity. There is a practical gap between "
        "these areas: thesis topic approval requires short-text semantic comparison, multilingual support, "
        "novelty scoring, and explainable warnings that supervisors can understand. This study addresses that "
        "gap by proposing a hybrid model specifically designed for thesis topic repositories.",
    )

    add_heading(doc, "3. Materials and Methods", 1)
    add_heading(doc, "3.1 Problem Formalisation", 2)
    add_para(
        doc,
        "Let a thesis topic record be represented as T = (title, abstract, keywords, domain, supervisor, year, "
        "language). Given a new topic q and a repository R = {T1, T2, ..., Tn}, the task is to estimate the "
        "novelty of q and classify it into one of three categories: duplicate, near duplicate, or novel. The "
        "model must also return the most similar historical topics and explain which components influenced the "
        "decision.",
    )
    add_para(
        doc,
        "The task differs from conventional document similarity because thesis titles are short, often "
        "multilingual, and domain-specific. In addition, the objective is not only to find the nearest topic, "
        "but also to evaluate whether the proposed topic contributes enough methodological, domain, or keyword "
        "novelty to justify approval.",
    )

    add_heading(doc, "3.2 Proposed Hybrid Architecture", 2)
    add_callout(
        doc,
        "Figure 1. Proposed model pipeline",
        "New topic -> preprocessing -> multilingual sentence embedding -> cosine similarity search -> "
        "keyword uniqueness -> domain rarity -> novelty score -> duplicate classification -> recommendation.",
    )
    add_para(
        doc,
        "The architecture contains six main stages. First, the topic text is normalised and tokenised. Second, "
        "the title, abstract, keywords, and domain are encoded into a dense vector using a multilingual "
        "sentence embedding model. Third, cosine similarity is computed between the query embedding and all "
        "stored topic embeddings. Fourth, keywords are analysed using corpus-level document frequency. Fifth, "
        "domain rarity is estimated from the frequency of domain-keyword combinations. Finally, the components "
        "are combined into a novelty score and mapped to a duplicate label.",
    )
    add_table(
        doc,
        "Table 1. Components of the proposed hybrid novelty model.",
        ["Component", "Purpose", "Output"],
        [
            [
                "Semantic similarity",
                "Measures meaning-level closeness between the new topic and existing topics.",
                "Maximum cosine similarity and top-k nearest topics",
            ],
            [
                "Keyword uniqueness",
                "Identifies whether the topic contains rare method, dataset, language, or domain terms.",
                "Normalised inverse document-frequency score",
            ],
            [
                "Domain rarity",
                "Estimates whether the domain and keyword combination is common in the repository.",
                "Domain-keyword rarity score",
            ],
            [
                "Novelty score",
                "Combines distance, keyword uniqueness, and domain rarity into one interpretable measure.",
                "Continuous value from 0 to 1",
            ],
            [
                "Duplicate classification",
                "Transforms similarity evidence into an operational warning for the platform.",
                "duplicate / near_duplicate / novel",
            ],
        ],
        [2200, 4600, 2560],
    )

    add_heading(doc, "3.3 Semantic Similarity", 2)
    add_para(
        doc,
        "For each topic, the model forms a text representation by combining the title, abstract, keywords, and "
        "domain. The title and keywords are treated as high-signal fields because they usually contain the "
        "main object, method, and application area of a thesis project. The resulting text is encoded as an "
        "embedding vector e. For E5-style models, query and passage prefixes are used to align inference with "
        "the model's retrieval-oriented training format.",
    )
    add_para(
        doc,
        "The semantic similarity between a query topic q and an existing topic Ti is computed using cosine "
        "similarity: cos(eq, ei) = (eq * ei) / (||eq|| ||ei||). The semantic similarity component of the "
        "model is the maximum similarity over the repository: SemSim(q) = max_i cos(eq, ei). A high value "
        "indicates that at least one existing topic is semantically close to the proposed topic.",
    )

    add_heading(doc, "3.4 Keyword Uniqueness", 2)
    add_para(
        doc,
        "Semantic similarity is not equivalent to novelty. Therefore, the model computes a keyword uniqueness "
        "component. Candidate keywords are taken from the user-provided keyword list and from extracted "
        "unigrams and bigrams in the title and abstract. For each keyword k, document frequency df(k) is "
        "computed over the repository. The uniqueness of a keyword is estimated through a normalised inverse "
        "document-frequency score. Unseen or rare keywords receive higher scores, while frequent generic "
        "terms receive lower scores.",
    )
    add_para(
        doc,
        "The keyword uniqueness score KWU(q) is the mean normalised uniqueness score of the candidate keywords "
        "in q. The model also returns rare and common keywords to make the decision interpretable for academic "
        "users.",
    )

    add_heading(doc, "3.5 Domain Rarity", 2)
    add_para(
        doc,
        "Domain rarity captures whether the proposed combination of domain and keywords is common in the "
        "institutional repository. If a domain appears frequently, a new topic in that domain is not "
        "automatically non-novel; however, a frequently repeated domain-keyword pair can indicate duplication. "
        "The domain rarity component combines two signals: the inverse frequency of the domain itself and the "
        "inverse frequency of each domain-keyword pair.",
    )
    add_para(
        doc,
        "The domain rarity score is defined as DR(q) = 0.4 * DomainInverseFrequency + 0.6 * "
        "MeanPairRarity. This weighting prioritises rare domain-keyword combinations while still accounting "
        "for heavily saturated domains.",
    )

    add_heading(doc, "3.6 Final Novelty Score and Classification", 2)
    add_para(
        doc,
        "The final novelty score is calculated as:",
    )
    add_callout(
        doc,
        "Novelty scoring formula",
        "NoveltyScore(q) = alpha * (1 - SemSim(q)) + beta * KWU(q) + gamma * DR(q), "
        "where alpha = 0.60, beta = 0.25, and gamma = 0.15 by default.",
    )
    add_para(
        doc,
        "The default weights reflect the assumption that semantic distance is the strongest signal, while "
        "keyword uniqueness and domain rarity provide additional evidence of research novelty. The weights "
        "can be tuned experimentally using expert-labelled validation pairs.",
    )
    add_para(
        doc,
        "Duplicate classification is based primarily on maximum semantic similarity. If SemSim(q) >= 0.90, "
        "the topic is classified as duplicate. If 0.75 <= SemSim(q) < 0.90, it is classified as near duplicate. "
        "Otherwise, it is classified as novel. These thresholds are initial values and should be optimised for "
        "the institutional dataset.",
    )

    add_heading(doc, "3.7 Implementation Architecture", 2)
    add_para(
        doc,
        "The prototype is implemented as a modular Python package. The core model can run with a deterministic "
        "hashing embedding backend for offline testing or with SentenceTransformer models for production-like "
        "semantic evaluation. A FastAPI service exposes the model through an /assess endpoint. PostgreSQL and "
        "pgvector are used in the deployment design to store topic metadata and embedding vectors in one "
        "database. For repositories with fewer than several thousand topics, exact cosine similarity is "
        "sufficient. For larger repositories, pgvector or FAISS-style nearest-neighbour indexing can be used.",
    )

    add_heading(doc, "4. Experimental Design", 1)
    add_para(
        doc,
        "Following the methodology of a result-oriented research article, the evaluation should prove not only "
        "that the model can compute similarity, but that the hybrid approach improves duplicate detection and "
        "novelty assessment compared with simpler baselines. The final experiment should use real anonymised "
        "institutional thesis topic records and expert-labelled topic pairs.",
    )
    add_table(
        doc,
        "Table 2. Recommended experimental protocol for the full study.",
        ["Element", "Recommended design"],
        [
            ["Dataset", "Anonymised historical diploma and master thesis topics from the university repository."],
            ["Minimum fields", "Title, abstract or description, keywords, domain, supervisor, year, language."],
            ["Labels", "Expert-labelled pairs: duplicate, near duplicate, novel."],
            ["Baselines", "TF-IDF cosine, BM25 or keyword overlap, Sentence-BERT-only, E5-only."],
            ["Proposed model", "Hybrid semantic similarity + keyword uniqueness + domain rarity."],
            ["Metrics", "Precision, recall, macro-F1, class-wise F1, confusion matrix, ROC-AUC if binary setup is used."],
            ["Validation", "Stratified split or k-fold validation over labelled pairs; threshold tuning on validation set."],
            ["Ablation", "Remove keyword uniqueness, remove domain rarity, compare with semantic-only scoring."],
        ],
        [2200, 7160],
    )
    add_para(
        doc,
        "The dataset should be split so that threshold tuning and final testing are separated. If the number "
        "of labelled pairs is small, k-fold cross-validation is preferable. Because near-duplicate cases are "
        "usually less frequent than novel pairs, macro-F1 and class-wise recall should be reported rather than "
        "accuracy alone.",
    )

    add_heading(doc, "5. Prototype Validation", 1)
    add_para(
        doc,
        "A preliminary prototype validation was conducted to verify the computational consistency of the "
        "pipeline. The pilot set contained 10 synthetic or manually prepared topic records and 6 labelled "
        "topic pairs. This test was not intended to prove generalisation; its purpose was to verify that the "
        "model pipeline, thresholds, and output schema behave correctly before institutional data collection.",
    )
    add_table(
        doc,
        "Table 3. Prototype sanity-check results on the small pilot pair set.",
        ["Class", "Precision", "Recall", "F1-score", "Support"],
        [
            ["duplicate", "1.000", "1.000", "1.000", "2"],
            ["near_duplicate", "1.000", "1.000", "1.000", "2"],
            ["novel", "1.000", "1.000", "1.000", "2"],
            ["macro average", "1.000", "1.000", "1.000", "6"],
        ],
        [2300, 1700, 1700, 1700, 1960],
    )
    add_para(
        doc,
        "These values must be interpreted cautiously. The pilot set is small and partly constructed to test "
        "edge cases. Therefore, the table demonstrates that the model implementation is operational, not that "
        "the model has achieved publication-ready performance. A full Scopus-oriented submission should "
        "replace this section with results from a larger, independently labelled institutional dataset.",
    )

    add_heading(doc, "6. Discussion", 1)
    add_para(
        doc,
        "The proposed hybrid model has several practical advantages for academic project management platforms. "
        "First, it provides a measurable warning when a proposed topic is too close to existing work. Second, "
        "it avoids the limitation of exact keyword search by using semantic embeddings. Third, it remains "
        "interpretable because the final decision is accompanied by similar topics, rare keywords, common "
        "keywords, and component-level scores. This is important in academic administration, where automated "
        "recommendations should support rather than replace supervisor judgement.",
    )
    add_para(
        doc,
        "The model is also suitable for multilingual repositories. In universities where thesis topics may be "
        "written in Russian, English, Kazakh, or mixed language, monolingual English embeddings can miss "
        "semantic similarity across languages. A multilingual E5 or multilingual MiniLM backend is therefore "
        "more appropriate than an English-only embedding model. The architecture also allows additional "
        "institution-specific components, such as supervisor workload, department research areas, or approved "
        "topic templates.",
    )
    add_para(
        doc,
        "From a research perspective, the main novelty is not the invention of a new transformer architecture. "
        "Rather, the contribution is the design and validation of an explainable hybrid pipeline for a concrete "
        "academic management problem. This aligns with applied machine learning research where novelty can be "
        "demonstrated through a reproducible pipeline, new feature combination, domain-specific evaluation, and "
        "clear comparison against baselines.",
    )

    add_heading(doc, "7. Limitations", 1)
    add_para(
        doc,
        "The current version has several limitations. First, the preliminary validation uses a very small pilot "
        "set and cannot support strong claims about generalisation. Second, threshold values such as 0.90 and "
        "0.75 are initial assumptions and must be tuned on real expert-labelled data. Third, keyword extraction "
        "is intentionally lightweight; domain-specific terminology dictionaries or supervised keyphrase "
        "extraction may improve performance. Fourth, the model focuses on topic novelty and duplicate risk, "
        "not on the scientific quality or feasibility of the proposed research plan. Finally, expert judgement "
        "remains necessary for borderline cases.",
    )

    add_heading(doc, "8. Conclusion", 1)
    add_para(
        doc,
        "This study proposed a hybrid semantic model for thesis topic novelty assessment and duplicate "
        "detection in academic project management platforms. The model combines multilingual sentence "
        "embeddings, cosine similarity, keyword uniqueness, domain rarity, novelty scoring, and threshold-based "
        "classification. The approach is reproducible, explainable, and compatible with practical deployment "
        "using FastAPI and PostgreSQL with pgvector. A pilot validation confirmed that the implementation "
        "operates as expected, while a full institutional experiment is required to establish empirical "
        "performance. Future work will focus on collecting expert-labelled topic pairs, comparing multilingual "
        "embedding backends, optimising thresholds, conducting ablation studies, and integrating the novelty "
        "module with supervisor assignment and milestone monitoring functions.",
    )

    add_heading(doc, "References", 1)
    references = [
        "A. Vaswani, N. Shazeer, N. Parmar, J. Uszkoreit, L. Jones, A. N. Gomez, L. Kaiser, and I. Polosukhin, "
        "\"Attention Is All You Need,\" Advances in Neural Information Processing Systems, 2017. "
        "https://arxiv.org/abs/1706.03762",
        "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, \"BERT: Pre-training of Deep Bidirectional "
        "Transformers for Language Understanding,\" NAACL-HLT, 2019. https://arxiv.org/abs/1810.04805",
        "N. Reimers and I. Gurevych, \"Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks,\" "
        "EMNLP-IJCNLP, 2019. https://arxiv.org/abs/1908.10084",
        "L. Wang, N. Yang, X. Huang, B. Jiao, L. Yang, D. Jiang, R. Majumder, and F. Wei, \"Text Embeddings by "
        "Weakly-Supervised Contrastive Pre-training,\" arXiv:2212.03533, 2022. https://arxiv.org/abs/2212.03533",
        "J. Johnson, M. Douze, and H. Jegou, \"Billion-scale similarity search with GPUs,\" IEEE Transactions "
        "on Big Data, vol. 7, no. 3, pp. 535-547, 2021. https://arxiv.org/abs/1702.08734",
        "S. Robertson and H. Zaragoza, \"The Probabilistic Relevance Framework: BM25 and Beyond,\" Foundations "
        "and Trends in Information Retrieval, vol. 3, no. 4, pp. 333-389, 2009. "
        "https://ir.webis.de/anthology/2009.ftir_journal-ir0anthology0volumeA3A4.0/",
        "F. Pedregosa et al., \"Scikit-learn: Machine Learning in Python,\" Journal of Machine Learning "
        "Research, vol. 12, pp. 2825-2830, 2011. https://www.jmlr.org/papers/v12/pedregosa11a.html",
        "pgvector, \"Open-source vector similarity search for Postgres,\" official repository/documentation. "
        "https://github.com/pgvector/pgvector",
        "Sentence Transformers, \"paraphrase-multilingual-MiniLM-L12-v2 model card,\" Hugging Face. "
        "https://huggingface.co/sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
    ]
    for idx, ref in enumerate(references, start=1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(ref)
        set_font(r, size=10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_article()
    print(out)
