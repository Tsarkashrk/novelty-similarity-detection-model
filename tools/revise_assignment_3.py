"""Revise the submitted Assignment 3 against the current public repository."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    "/Users/tsarevich/Downloads/"
    "Assignment_3_Software_Development_and_Integration_Alisher_Tokesh.docx"
)
OUTPUT = ROOT / "outputs/Assignment_3_Updated_Alisher_Tokesh.docx"


def replace(paragraph, value: str) -> None:
    paragraph.clear()
    paragraph.add_run(value)


def add_link(paragraph, label: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "215477")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_body(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def insert_paragraph(document, before, value: str, style: str = "Normal"):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(value)
    if style == "Heading 1":
        properties = paragraph._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number_id = OxmlElement("w:numId")
        number_id.set(qn("w:val"), "0")
        numbering.extend([level, number_id])
        properties.append(numbering)
    before.addprevious(paragraph._element)
    if style == "Normal":
        set_body(paragraph)
    return paragraph


def insert_table(document, before, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = 1
    for index, header in enumerate(headers):
        table.columns[index].width = Cm(widths[index])
        table.rows[0].cells[index].width = Cm(widths[index])
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Cm(widths[index])
            cells[index].text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
                    if row_index == 0:
                        run.bold = True
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{side}")
                element.set(qn("w:w"), "100")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tc_pr.append(margins)
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:tblHeader"))
    before.addprevious(table._element)
    return table


def main() -> None:
    document = Document(SOURCE)
    paragraphs = list(document.paragraphs)

    changes = {
        5: (
            "This report continues the thesis-topic novelty screening project in Assignments 1 and 2. "
            "Assignment 1 set a 17-week plan and a USD 43,916 budget, including contingency. "
            "Assignment 2 selected stage-gated iterative R&D and a modular monolith with replaceable adapters. "
            "Assignment 3 documents the public scientific software module, Git and GitHub workflow, tests, "
            "and CI. The latest repository increment adds an OpenAlex-derived silver-labelled corpus and "
            "makes empirical error analysis possible."
        ),
        6: (
            "The module assesses a proposed title, abstract, keywords, domain, and language using semantic "
            "similarity, keyword uniqueness, and domain rarity. It still retrieves candidates in memory; "
            "pgvector remains a planned integration. On the new 130-pair silver benchmark, the dependency-free "
            "hashing backend finds only 6 of 40 near-duplicate pairs. This is preliminary evidence of a "
            "quality gap, not a deployment-ready result. A high-similarity flag must prompt human review, "
            "never automatic academic rejection."
        ),
        9: (
            "The development process moves from requirements and data governance through design, "
            "implementation, verification, and controlled release. It remains iterative because corpus "
            "composition, multilingual embeddings, and thresholds must be tested against examples. Stage "
            "gates protect student-facing decisions: approved access, benchmark integrity, model quality, "
            "security, and pilot readiness all need explicit evidence. The OpenAlex corpus completes a "
            "prototype data increment, not the institutional data-approval gate."
        ),
        16: (
            "The pyproject.toml file keeps the core package dependency-light and defines API, embeddings, "
            "experiment, and development extras. A deterministic hashing backend makes CI fast; "
            "SentenceTransformers is optional for a research environment with a pinned multilingual model. "
            "The new OpenAlex builder uses Python's standard-library HTTP, JSON, and CSV tools, writes raw "
            "response snapshots, and generates silver labels from controlled variants. Its data-provenance "
            "benefit must be balanced against label bias and the need for expert review."
        ),
        17: (
            "For every model comparison, record the Git commit, corpus revision, language mix, model ID, "
            "weights, thresholds, and evaluation script. Assignment 2 set pilot targets of duplicate F1 "
            "at least 0.82, near-duplicate recall at least 0.75, false high-risk rate at most 8 percent, "
            "and top-k latency at most 1.5 seconds. The new pairwise hashing result does not meet the "
            "near-duplicate recall target; none of these pilot KPIs can be certified from the current "
            "synthetic-heavy benchmark."
        ),
        21: (
            "The public repository uses a main branch and a GitHub remote. A normal change starts with "
            "a focused issue, uses a short-lived branch, and carries a test or reproducible example into "
            "review. The README describes setup and the CI badge; the MIT license states reuse terms. "
            "Bug and research-task templates request evidence while warning against posting private "
            "student data. The dataset is versioned with the source so its records can be tied to a "
            "specific commit."
        ),
        22: (
            "Commit 897960e introduced the workflow, license, issue templates, README update, and CLI "
            "test. Commit 6d7a876 added the OpenAlex silver corpus, raw snapshots, and builder. Issue 1 "
            "continues to track the pgvector adapter with acceptance criteria. A tag for a validated "
            "release should include the model configuration, benchmark revision, and a human approval "
            "record; the current commit is not such a release."
        ),
        26: (
            "The workflow at .github/workflows/ci.yml runs tests on Python 3.10 and 3.12, compiles the "
            "source, then builds and uploads a wheel after successful tests on main. It has read-only "
            "repository permissions and no deployment step. This is CI plus a delivery candidate, not "
            "continuous deployment. The workflow currently does not validate the OpenAlex CSV schema, "
            "execute evaluate_thresholds.py, or enforce a model-quality threshold."
        ),
        27: (
            "GitHub Actions run 37896238736 completed successfully for commit 6d7a876, confirming that "
            "the package tests and build still pass after the corpus addition. The two test jobs and "
            "package job succeeded. Locally, the three existing tests pass with PYTHONPATH=src. "
            "Neither result demonstrates multilingual quality or near-duplicate recall; those require "
            "separate evaluation gates. Automatic publication or university deployment is not configured."
        ),
        29: (
            "Scientific software must make a claim reproducible and falsifiable. A novelty score depends "
            "on the reference corpus, preprocessing, embedding backend, weights, thresholds, and label "
            "protocol. The new corpus improves traceability through OpenAlex identifiers, raw responses, "
            "and a fixed generation seed, but its labels are generated rather than adjudicated by "
            "academic experts. Engineering tests and scientific validation answer different questions."
        ),
        31: (
            "The core package uses standard-library dataclasses and deterministic hashing embeddings. "
            "FastAPI/Pydantic and SentenceTransformers are optional dependencies; pandas, scikit-learn, "
            "and matplotlib support experiments. The OpenAlex builder needs no additional runtime "
            "library. A research run should pin the package and embedding model, preserve the raw API "
            "snapshot and corpus hash, and avoid public student records. API contract and dependency "
            "behavior still need a dedicated integration test before the pilot."
        ),
        33: (
            "The existing tests cover duplicate and unrelated-topic model behavior plus a CLI JSON "
            "smoke check. The pairwise evaluation script now runs on the OpenAlex files, but is not part "
            "of CI and assesses each query against one nominated candidate, not full-corpus retrieval. "
            "Next tests should validate dataset IDs and labels, run an API contract check, test pgvector "
            "migration and restart behavior, and measure retrieval plus scoring on a pilot-sized corpus. "
            "Language-stratified metrics require separate Russian and Kazakh data."
        ),
        34: (
            "The principal scientific risk is that passing software checks can create a false sense of "
            "model validity. On the silver benchmark, 34 of 40 near-duplicate pairs are classified as "
            "novel by the hashing backend. Exact duplicate variants are easy by construction, while "
            "distant-domain novel pairs are also easier than realistic borderline cases. Mitigation is "
            "expert review of a stratified subset, genuine hard negatives, a frozen test split by source "
            "topic family, multilingual encoder comparison, and manual review of every high-risk flag."
        ),
        36: (
            "Each experiment should record the code commit, dependency and model versions, corpus hash, "
            "query/candidate split, language distribution, thresholds, and exact command. Seed 42 "
            "stabilizes pair sampling for an unchanged input list, and raw JSON preserves an audit trail. "
            "The current builder still fetches the live API rather than replaying those snapshots; "
            "offline regeneration is a needed improvement. Source-derived variants must stay in "
            "their original's data split to avoid leakage."
        ),
        38: (
            "The practical module is the thesis_novelty package. A CSV loader creates topic records; "
            "preprocessing and an interchangeable embedding backend encode the query and corpus. "
            "The model ranks candidates in memory by cosine similarity, then combines semantic "
            "distance, keyword uniqueness, and domain rarity with weights 0.60, 0.25, and 0.15. "
            "Default similarity cutoffs are 0.90 for duplicate and 0.75 for near_duplicate. The CLI "
            "and FastAPI factory expose the assessment; the new data files supply a larger evaluation "
            "fixture, not a new production retrieval service."
        ),
        40: (
            "The public repository includes source code, tests, sample fixtures, README, MIT license, "
            "issue templates, CI workflow, a pgvector schema, and the OpenAlex silver corpus with raw "
            "snapshots and a builder. The latest public Actions run verifies the committed package "
            "checks and wheel build. The corpus and evaluation script make the scientific limitations "
            "inspectable."
        ),
        41: (
            "The references link the repository, CI, dataset, and evaluation code. Issue 1 records "
            "the next storage integration; all evidence is public and uses no private student topics."
        ),
        42: (
            "This remains a delivery candidate: pgvector, approved institutional data, multilingual "
            "quality, latency, and deployment are open. The OpenAlex set bridges the small sample "
            "fixture and a future expert-labelled university benchmark, but cannot replace it."
        ),
        44: (
            "The methodology and modular architecture from Assignments 1 and 2 still fit this small "
            "scientific service. The public Git history, issue workflow, dependency-light package, "
            "automated tests, and successful CI run demonstrate development integration. The OpenAlex "
            "increment materially improves reproducible evaluation, yet the measured near-duplicate "
            "recall of 0.15 exposes a gap against the 0.75 pilot target. Therefore this commit supports "
            "further R&D, not student-facing deployment."
        ),
        45: (
            "The main implementation challenge is the distinction between software correctness and "
            "scientific validity. Three smoke tests and a successful wheel build coexist with a poor "
            "near-duplicate result. The current silver labels are useful for diagnosis, but 40 near "
            "pairs omit the abstract, 131 of 160 real records are English, and the set has no Russian "
            "or Kazakh records. The evaluation is pairwise, not end-to-end retrieval. These facts "
            "limit any claim about university multilingual performance."
        ),
        46: (
            "Retain the 17-week and USD 43,916 baselines, but use the quality gate to prioritize "
            "expert-labelled Russian and Kazakh examples, hard negatives, family-disjoint splits, "
            "multilingual embedding comparison, and full-corpus retrieval tests. Add dataset integrity "
            "checks to CI, then evaluate frozen model metrics in a separately provisioned job. "
            "Do not release until quality, security, latency, and human-review gates pass."
        ),
    }
    for index, value in changes.items():
        replace(paragraphs[index], value)

    replace(paragraphs[43], "8. Conclusions and recommendations")
    replace(paragraphs[7], "Public evidence: ")
    add_link(
        paragraphs[7],
        "repository",
        "https://github.com/Tsarkashrk/novelty-similarity-detection-model",
    )
    paragraphs[7].add_run("; ")
    add_link(
        paragraphs[7],
        "successful CI run for 6d7a876",
        "https://github.com/Tsarkashrk/novelty-similarity-detection-model/actions/runs/37896238736",
    )
    paragraphs[7].add_run(".")

    technology = document.tables[1]
    row = technology.add_row().cells
    row[0].text = "OpenAlex API and versioned snapshots"
    row[1].text = "Public dissertation metadata and reproducible silver-pair construction"
    row[2].text = "Labels are synthetic or heuristic; verify source terms and obtain expert labels"
    for cell in row:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    components = document.tables[4]
    row = components.add_row().cells
    row[0].text = "OpenAlex builder and CSVs"
    row[1].text = "Versioned research corpus and pairwise benchmark"
    row[2].text = "Keep data generation separate from assessment and release gates"
    for cell in row:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    if "Title" not in document.styles:
        title_style = document.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title_style = document.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.line_spacing = 1.0
    title_style.paragraph_format.space_after = Pt(12)
    paragraphs[1].style = title_style
    paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[3].paragraph_format.space_after = Pt(80)
    for index, paragraph in enumerate(document.paragraphs):
        if index > 3 and paragraph.style.name == "Normal":
            set_body(paragraph)
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)

    # Preserve the source's A4 page and restrained report layout.
    section = document.sections[0]
    section.bottom_margin = Cm(2.3)

    anchor = paragraphs[43]._element
    insert_paragraph(document, anchor, "7. Updated OpenAlex corpus and evaluation", "Heading 1")
    insert_paragraph(
        document,
        anchor,
        "The public OpenAlex increment contains 160 dissertation records and 70 controlled "
        "variants: 230 topics and 130 silver-labelled pairs. Ten research themes, versioned "
        "raw responses, and source identifiers support auditing without private student work. "
        "The labels are rule-generated, not expert-reviewed.",
    )
    insert_table(
        document,
        anchor,
        ["Item", "Observed value", "Implication"],
        [
            ["Real records", "160; 131 English", "No Russian or Kazakh evaluation yet"],
            ["Derived variants", "30 exact, 40 paraphrased", "Easy positives and family leakage risk"],
            ["Novel pairs", "60 distant-domain pairs", "Not representative hard negatives"],
            ["Near-pair fields", "All 40 query abstracts empty", "Title/field mismatch affects scores"],
        ],
        [3.2, 4.0, 9.1],
    )
    caption = insert_paragraph(document, anchor, "Table 6 Corpus composition and validity limits", "Heading 3")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = False
    for run in caption.runs:
        run.bold = False
        run.italic = True
    insert_paragraph(
        document,
        anchor,
        "The builder labels exact copies as duplicate, controlled title paraphrases as "
        "near_duplicate, and cross-domain low-overlap pairs as novel. It selects negative pairs "
        "with title/keyword Jaccard overlap at most 0.08. Because this construction defines the "
        "ground truth, high scores on exact copies or distant negatives would not alone establish "
        "real-world discrimination. The near-duplicate variants omit query abstracts, while their "
        "source records may retain them; this makes the assessment sensitive to field availability.",
    )
    insert_table(
        document,
        anchor,
        ["Actual label", "Predicted duplicate", "Predicted near", "Predicted novel", "Recall"],
        [
            ["Duplicate", "30", "0", "0", "1.00"],
            ["Near-duplicate", "0", "6", "34", "0.15"],
            ["Novel", "0", "0", "60", "1.00"],
        ],
        [3.4, 3.3, 3.1, 3.2, 3.3],
    )
    caption = insert_paragraph(document, anchor, "Table 7 Pairwise hashing benchmark at default thresholds", "Heading 3")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = False
    for run in caption.runs:
        run.bold = False
        run.italic = True
    insert_paragraph(
        document,
        anchor,
        "I ran evaluate_thresholds.py against the committed OpenAlex CSVs with the default hashing "
        "backend and thresholds 0.90/0.75. Duplicate F1 was 1.00, near-duplicate F1 was 0.261, "
        "novel F1 was 0.779, and macro F1 was 0.680. The near-duplicate recall of 0.15 is far below "
        "Assignment 2's 0.75 pilot target. The 60 silver novel pairs produced no false high-risk "
        "flag, but this narrow, easy-negative sample cannot validate the proposed 8 percent ceiling "
        "on real university topics. The script compares each query with one nominated candidate; "
        "top-k retrieval accuracy and latency remain unmeasured.",
    )
    insert_paragraph(
        document,
        anchor,
        "Recommended next increment. Freeze a family-disjoint expert-reviewed test set, add Russian "
        "and Kazakh topics plus borderline hard negatives, compare hashing with a pinned multilingual "
        "encoder, and report confidence intervals and per-language errors. Add fast dataset-schema "
        "and reference-integrity checks to ordinary CI. Run model evaluation in a separately "
        "provisioned workflow with cached weights and pinned dependencies, then require a reviewed "
        "metric gate before manual release. The builder also needs verified TLS and an offline "
        "snapshot-replay mode. This keeps software feedback fast without confusing a passing build "
        "with scientific acceptance.",
    )

    references_heading = document.add_paragraph("9. References and public evidence", style="Heading 1")
    properties = references_heading._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "0")
    numbering.extend([level, number_id])
    properties.append(numbering)
    for label, url in [
        ("Project repository", "https://github.com/Tsarkashrk/novelty-similarity-detection-model"),
        ("CI configuration", "https://github.com/Tsarkashrk/novelty-similarity-detection-model/blob/main/.github/workflows/ci.yml"),
        ("Successful run for commit 6d7a876", "https://github.com/Tsarkashrk/novelty-similarity-detection-model/actions/runs/37896238736"),
        ("OpenAlex dataset documentation", "https://github.com/Tsarkashrk/novelty-similarity-detection-model/blob/main/data/openalex_silver/README.md"),
        ("Dataset builder", "https://github.com/Tsarkashrk/novelty-similarity-detection-model/blob/main/tools/build_openalex_silver_dataset.py"),
        ("Evaluation script", "https://github.com/Tsarkashrk/novelty-similarity-detection-model/blob/main/experiments/evaluate_thresholds.py"),
    ]:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.add_run(f"{label}: ")
        add_link(paragraph, url, url)
        set_body(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn("w:cantSplit")) is None:
                properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0 and properties.find(qn("w:tblHeader")) is None:
                properties.append(OxmlElement("w:tblHeader"))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
