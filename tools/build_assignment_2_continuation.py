from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path('/Users/tsarevich/Downloads/Assignment 1 Alisher Tokesh.docx')
OUT = ROOT / 'outputs' / 'Assignment_2_Continuation_Alisher_Tokesh.docx'


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_para(paragraph, before=0, after=8, line=1.08, first_indent=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_indent:
        fmt.first_line_indent = Inches(0.25)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def add_text(doc, text, *, first_indent=True, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_para(p, after=after, first_indent=first_indent)
    run = p.add_run(text)
    set_font(run)
    return p


def add_lead(doc, lead, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_para(p, after=8, first_indent=True)
    r = p.add_run(lead)
    set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para(p, before=12 if level == 1 else 8, after=5, line=1.0)
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 14, bold=True)
    return p


def cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '6')
        node.set(qn('w:color'), '000000')


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(int(width * 1440)))
    tc_w.set(qn('w:type'), 'dxa')


def table(doc, headers, rows, widths, font_size=9.4):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_width(c, widths[i])
        cell_border(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.0)
        r = p.add_run(text)
        set_font(r, size=font_size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            c = cells[i]
            set_cell_width(c, widths[i])
            cell_border(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0, line=1.0)
            r = p.add_run(str(text))
            set_font(r, size=font_size)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, before=4, after=7, line=1.0)
    r = p.add_run(text)
    set_font(r, size=10, italic=True)


def diagram(doc):
    labels = [
        'Coordinator or reviewer', 'FastAPI adapter', 'Assessment service', 'Hybrid scoring policy',
        'Reviewer decision and audit', 'Topic repository port', 'Embedding provider port', 'Configuration registry',
        'PostgreSQL with pgvector', 'Multilingual embedding model', 'Metrics and quality dashboard', 'Background evaluation jobs',
    ]
    t = doc.add_table(rows=3, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for index, label in enumerate(labels):
        cell = t.cell(index // 4, index % 4)
        set_cell_width(cell, 1.6)
        cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.0)
        r = p.add_run(label)
        set_font(r, size=9, bold=index < 4)
    caption(doc, 'Figure 1. Target modular monolith with adapters and operational boundaries.')


def page_break(doc):
    doc.add_page_break()


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(REFERENCE)
    clear_body(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Page 1
    doc.add_paragraph().paragraph_format.space_after = Pt(60)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=22, line=1.1)
    set_font(p.add_run('Development Methodology and Architecture Assessment for a Novelty Similarity Detection Model'), size=20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=36)
    set_font(p.add_run('Alisher Tokesh | CSE-2507M'), size=16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=90)
    set_font(p.add_run('Assignment 2'), size=16)
    add_text(doc, 'This report continues Assignment 1 and assesses whether the recommended development methodology and architecture fit the same novelty similarity detection pilot. The analysis keeps the 17-week critical path, USD 43,916 budget, human review policy, and quality targets established in the preceding project plan.', first_indent=False, after=14)
    add_text(doc, 'Main conclusion. A stage-gated iterative R&D approach and a modular monolith with explicit adapters are appropriate for the pilot. They preserve the ability to calibrate the multilingual model and validate reviewers feedback, while preventing budget and integration complexity from outrunning the evidence needed for a production decision.', first_indent=False, after=0)
    page_break(doc)

    # Page 2
    add_heading(doc, 'Table of Contents')
    for line in [
        '1 Project Analysis ........................................ 3',
        '1.1 Goals Requirements and Context ........................ 3',
        '1.2 Constraints Risks and Priorities ...................... 4',
        '2 Development Methodology Assessment ...................... 5',
        '2.1 Selected Method and Fit ............................... 5',
        '2.2 Method Risks and Alternatives .......................... 6',
        '3 Architectural Approach Assessment ........................ 7',
        '3.1 Architectural Style and Alignment ..................... 7',
        '3.2 Quality Attributes and Adjustments .................... 8',
        '4 Conclusions and Recommendations .......................... 9',
        'References ................................................ 9',
    ]:
        p = doc.add_paragraph()
        set_para(p, after=4)
        set_font(p.add_run(line), size=11)
    add_heading(doc, 'Continuation from Assignment 1', 2)
    add_text(doc, 'Assignment 1 defined the novelty model as a controlled decision-support service rather than an autonomous approval system. Its scope includes a cleaned historical corpus, threshold calibration, FastAPI with pgvector search, a coordinator workflow, quality monitoring, pilot training, and closure. The present report evaluates the development decisions required to deliver those outputs within the approved baseline.', first_indent=False)
    add_lead(doc, 'Project baseline. ', 'The planned implementation lasts 17 weeks. The direct cost is USD 38,188 with a 15 percent contingency reserve of USD 5,728, giving a total budget of USD 43,916. These values are treated as constraints for methodology and architecture decisions, not merely as financial background.')
    page_break(doc)

    # Page 3
    add_heading(doc, '1 Project Analysis')
    add_heading(doc, '1.1 Goals Requirements and Context', 2)
    add_text(doc, 'The project aims to identify duplicate and near-duplicate thesis topics before final academic approval. It accepts a title, abstract, keywords, domain and language, compares the proposal with historical topics, and returns a novelty score, similarity evidence, and a decision label. Reviewers remain responsible for the final approval, especially when a topic is classified as high risk.', first_indent=False)
    table(doc, ['Category', 'Priority', 'Requirement', 'Evidence of completion'], [
        ['Functional', 'Must', 'Load a cleaned historical corpus and calculate top-k similar topics.', 'Versioned corpus and reproducible search result.'],
        ['Functional', 'Must', 'Return score components, label and similar topics through FastAPI.', 'API contract and reviewer-readable explanation.'],
        ['Functional', 'Should', 'Support reviewer confirmation or override with a reason.', 'Auditable decision record.'],
        ['Functional', 'Should', 'Provide a coordinator workflow and quality dashboard.', 'Pilot users can inspect cases and metrics.'],
        ['Nonfunctional', 'Must', 'Protect topic data and avoid automatic rejection.', 'Role-based access and human-in-the-loop policy.'],
        ['Nonfunctional', 'Must', 'Produce reproducible calibration and evaluation results.', 'Dataset, model and threshold versions are recorded.'],
        ['Nonfunctional', 'Should', 'Meet response and quality targets under pilot load.', 'KPI report and load-test evidence.'],
    ], [0.9, 0.7, 2.8, 2.1], font_size=8.7)
    caption(doc, 'Table 1. Prioritized requirements derived from the Assignment 1 scope.')
    page_break(doc)

    # Page 4
    add_heading(doc, '1.2 Constraints Risks and Priorities', 2)
    add_text(doc, 'The project is an applied ML and integration pilot, so its principal uncertainty is not only engineering effort. Historical data may be incomplete, language coverage may be uneven, and a technically high similarity score may still require contextual academic judgement. The schedule gives only two weeks of slack to corpus preparation and hybrid calibration; the API and workflow integration stream is critical.', first_indent=False)
    table(doc, ['Constraint or risk', 'Effect on project', 'Required control'], [
        ['17-week critical path', 'A delay in approvals, backend integration or pilot readiness delays closure.', 'Stage gates and weekly critical-path review.'],
        ['Budget USD 43,916', 'Premature service separation or excess experimentation consumes the reserve.', 'Release reserve only by approved change request.'],
        ['Inconsistent historical data', 'Noisy similarity and weak keyword statistics reduce trust.', 'Data audit, minimum fields and sampled label review.'],
        ['Kazakh Russian English coverage', 'Uneven language quality can create unfair results.', 'Language-tagged validation sets and multilingual embeddings.'],
        ['False high-risk flags', 'Students can lose trust if a flag is treated as a decision.', 'Mandatory reviewer review, appeal path and threshold tuning.'],
        ['Privacy and access controls', 'Approval delays can block model work and pilot use.', 'Early data-owner approval, restricted logging and RBAC.'],
    ], [1.5, 2.4, 2.6], font_size=8.8)
    caption(doc, 'Table 2. Project constraints and risk controls inherited from Assignment 1.')
    add_heading(doc, 'Pilot Quality Targets', 2)
    add_text(doc, 'The architecture and process are accepted only if they support duplicate-detection F1 of at least 0.82, near-duplicate recall of at least 0.75, top-k latency of no more than 1.5 seconds, at least 25 percent review-time saving, user acceptance of at least 4.0 out of 5, and a false high-risk rate no greater than 8 percent.', first_indent=False)
    page_break(doc)

    # Page 5
    add_heading(doc, '2 Development Methodology Assessment')
    add_heading(doc, '2.1 Selected Method and Fit', 2)
    add_text(doc, 'The proposed method is stage-gated iterative R&D delivery: short Scrum-style planning and review cycles for software and model work, a Kanban board with work-in-progress limits for data remediation, and formal gates at points where evidence or funding must be released. It is not a full bureaucratic waterfall. The gates protect the schedule, budget and academic-risk controls that Assignment 1 already established.', first_indent=False)
    table(doc, ['Period', 'Increment and decision gate', 'Primary evidence'], [
        ['Weeks 1 to 3', 'Scope, data access and baseline preparation', 'Data approval, duplicate policy, labelled-pair protocol.'],
        ['Weeks 3 to 6', 'Baseline model and corpus work', 'Reproducible hashing baseline and corpus quality report.'],
        ['Weeks 5 to 11', 'Multilingual embedding, pgvector and API integration', 'Top-k retrieval, service contract and versioned configuration.'],
        ['Weeks 11 to 13', 'Quality security and performance validation', 'F1, recall, latency, security and error analysis.'],
        ['Weeks 13 to 15', 'Coordinator pilot', 'Reviewer decisions, disagreement analysis and acceptance feedback.'],
        ['Weeks 15 to 17', 'Deployment training and closure', 'Rollout decision, training record and lessons learned.'],
    ], [1.1, 2.7, 2.7], font_size=8.8)
    caption(doc, 'Table 3. Iterative delivery aligned with the 17-week CPM baseline.')
    add_lead(doc, 'Why the fit is strong. ', 'Model thresholds, corpus quality and reviewer expectations cannot be fixed with confidence at the start. Iterations expose those uncertainties early; gates stop the team from treating a working API as proof that the decision-support service is trustworthy.')
    page_break(doc)

    # Page 6
    add_heading(doc, '2.2 Method Risks and Alternatives', 2)
    table(doc, ['Criterion', 'Stage-gated iterative R&D', 'Waterfall', 'Kanban only'], [
        ['Data and model uncertainty', 'High fit: baseline and calibration can change from evidence.', 'Low fit: feedback arrives after detailed commitments.', 'Medium fit: flexible but needs quality gates.'],
        ['17-week deadline', 'High fit: time boxes expose delays on critical activities.', 'Medium fit: predictable but inflexible around evidence.', 'Medium fit: flow may hide fixed milestones.'],
        ['Budget reserve control', 'High fit: funding and reserve link to gate evidence.', 'High fit: approval controls are simple.', 'Low fit unless added financial controls exist.'],
        ['Reviewer involvement', 'High fit: regular demos and label reviews.', 'Low fit: reviewers see the product late.', 'Medium fit: possible but not time-boxed.'],
        ['Operations after pilot', 'Medium fit: switch to Kanban after release.', 'Low fit: change requests become slow.', 'High fit: suitable for maintenance.'],
    ], [1.25, 2.0, 1.6, 1.6], font_size=8.4)
    caption(doc, 'Table 4. Comparative methodology assessment.')
    add_heading(doc, 'Methodology Risks and Responses', 2)
    add_text(doc, 'The chosen method can fail if experimentation has no stopping rule, if reviewers miss every review, or if the data-remediation queue grows without a limit. The response is to freeze a small benchmark before calibration, assign a named academic decision delegate, apply Kanban WIP limits, and treat a change that exceeds one iteration as a scoped research decision. A limited waterfall phase is justified only for privacy approval and procurement documents whose acceptance criteria are stable.', first_indent=False)
    add_lead(doc, 'Recommendation. ', 'Use Kanban as the operational control within the iterative method, not as a replacement for sprint reviews, pilot gates and budget authority.')
    page_break(doc)

    # Page 7
    add_heading(doc, '3 Architectural Approach Assessment')
    add_heading(doc, '3.1 Architectural Style and Alignment', 2)
    add_text(doc, 'The existing repository is already organized like a modular monolith: preprocessing, embeddings, keyword analysis, configuration, schemas, scoring, CSV I/O, CLI and FastAPI are separated. The EmbeddingBackend abstraction supports both a deterministic hashing fallback and a SentenceTransformer backend. The present prototype ranks in process memory, while the Assignment 1 plan reserves pgvector for production-style similarity search.', first_indent=False)
    add_text(doc, 'The recommended target remains a modular monolith with hexagonal boundaries. The assessment application service owns the use case; the hybrid scorer and policy remain independent of infrastructure; adapters connect the service to FastAPI, the topic repository, the embedding provider, audit logging, metrics and background evaluation. This is more proportionate than microservices for the 17-week, USD 43,916 pilot.', first_indent=False)
    diagram(doc)
    page_break(doc)

    # Page 8
    add_heading(doc, '3.2 Quality Attributes and Adjustments', 2)
    table(doc, ['Attribute', 'Current support', 'Required adjustment'], [
        ['Reliability', 'Unit tests cover duplicate and unrelated cases; health endpoint exists.', 'Add integration tests for storage, restart recovery and malformed corpus data.'],
        ['Performance', 'In-memory search is simple; pgvector HNSW schema is planned.', 'Use indexed pgvector retrieval, batch embedding and measured p95 latency.'],
        ['Security and privacy', 'No application authentication or audit adapter is visible.', 'Add platform RBAC, TLS, least-privilege DB role and redacted logs.'],
        ['Explainability', 'Score components and similar topics are returned by the model.', 'Persist configuration, candidates and reviewer rationale with each assessment.'],
        ['Maintainability', 'Modules and embedding abstraction isolate many changes.', 'Formalize repository, audit and provider ports; record ADRs and API versions.'],
        ['Reproducibility', 'Configurable weights, thresholds and deterministic fallback exist.', 'Version corpus, embedding model, threshold report and benchmark together.'],
    ], [1.15, 2.7, 2.7], font_size=8.4)
    caption(doc, 'Table 5. Architecture alignment with nonfunctional requirements and KPI targets.')
    add_heading(doc, 'Justified Architectural Adjustments', 2)
    add_text(doc, 'First, replace service-owned topic lists with a TopicRepository port and a pgvector implementation, while retaining scoring in the domain layer. Second, make the audit record a first-class adapter that stores corpus version, model version, thresholds, score components, returned candidates and reviewer action. Third, move ingestion, batch embedding and evaluation into controlled background jobs so long-running work cannot delay the assessment endpoint. Microservice extraction should wait until independent scaling or a separate security boundary produces a measured benefit.', first_indent=False)
    page_break(doc)

    # Page 9
    add_heading(doc, '4 Conclusions and Recommendations')
    add_text(doc, 'The selected methodology and architecture fit the same project baseline defined in Assignment 1. The 17-week critical path and USD 43,916 limit require a small-team design with visible evidence gates. At the same time, the multilingual novelty model cannot be safely completed through a rigid sequential plan because data quality, threshold calibration and reviewer interpretation are uncertain. Stage-gated iterative R&D therefore balances flexibility with funding and academic governance.', first_indent=False)
    add_text(doc, 'A modular monolith with hexagonal adapters is the appropriate architectural choice. It uses the existing repository boundaries, enables a pgvector transition without coupling the scoring policy to a database, and preserves a single deployable service for a constrained pilot. The important future investments are persistent corpus management, human-decision auditability, access control, benchmark governance and operational monitoring, rather than premature microservices.', first_indent=False)
    add_heading(doc, 'Prioritized Recommendations', 2)
    table(doc, ['Timing', 'Recommendation', 'Decision criterion'], [
        ['Weeks 1 to 3', 'Approve policy, access, benchmark and error definitions.', 'No model calibration before data and review rules are accepted.'],
        ['Weeks 3 to 11', 'Deliver a vertical slice with multilingual embeddings, pgvector and API.', 'Top-k evidence is reproducible and below the latency target.'],
        ['Weeks 11 to 15', 'Add audit trail, RBAC, test evidence and reviewer pilot.', 'F1, recall, false-high-risk rate and user acceptance meet targets.'],
        ['Weeks 15 to 17', 'Deploy only after the pilot gate; otherwise use reserve for correction.', 'Production decision is evidence-based, not deadline-based.'],
    ], [1.1, 3.3, 2.15], font_size=8.7)
    caption(doc, 'Table 6. Recommendations linked to the original schedule and quality gates.')
    add_heading(doc, 'References', 2)
    for ref in [
        '1. Assignment 1 Alisher Tokesh. Project Management and Financial Justification for a Novelty Similarity Detection Model. Course report.',
        '2. Hybrid Thesis Topic Novelty Assessment Model. README.md and source modules in novelty-similarity-detection-model.',
        '3. Assignment 2 brief supplied by the course. Methodology and architectural approach assessment.',
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_para(p, after=3, line=1.0)
        set_font(p.add_run(ref), size=10)

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
