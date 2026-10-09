from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs"
ASSET_DIR = OUTPUT_DIR / "assignment_1_assets"
DOCX_PATH = OUTPUT_DIR / "Assignment_1_Project_Management_Financial_Justification_Novelty_Model.docx"


BLACK = RGBColor(0, 0, 0)
DARK = "#111827"
MUTED = "#4B5563"
BLUE = "#1D4ED8"
BLUE_DARK = "#1E3A8A"
BLUE_LIGHT = "#DBEAFE"
GRAY = "#6B7280"
GRAY_LIGHT = "#F3F4F6"
GREEN = "#15803D"
GREEN_LIGHT = "#DCFCE7"
AMBER = "#B45309"
AMBER_LIGHT = "#FEF3C7"
RED = "#B91C1C"
RED_LIGHT = "#FEE2E2"
BORDER = "#D9D9D9"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "",
        "/Library/Fonts/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size=size)
            except OSError:
                pass
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.ImageFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def draw_centered(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    fnt: ImageFont.ImageFont,
    fill: str = DARK,
) -> None:
    w, h = text_size(draw, text, fnt)
    draw.text((xy[0] - w / 2, xy[1] - h / 2), text, font=fnt, fill=fill)


def wrap_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    fnt: ImageFont.ImageFont,
    width: int,
) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if text_size(draw, candidate, fnt)[0] <= width or not current:
                current = candidate
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines or [""]


def draw_wrapped_centered(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    text: str,
    fnt: ImageFont.ImageFont,
    fill: str = DARK,
    line_gap: int = 6,
) -> None:
    x1, y1, x2, y2 = rect
    lines = wrap_text(draw, text, fnt, x2 - x1 - 24)
    heights = [text_size(draw, line, fnt)[1] for line in lines]
    total_h = sum(heights) + max(0, len(lines) - 1) * line_gap
    y = y1 + (y2 - y1 - total_h) / 2
    for line, h in zip(lines, heights):
        w, _ = text_size(draw, line, fnt)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill=fill)
        y += h + line_gap


def rounded_box(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    fill: str,
    outline: str = BORDER,
    radius: int = 22,
    width: int = 3,
) -> None:
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    fill: str = GRAY,
    width: int = 4,
) -> None:
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    head_len = 18
    for delta in (math.pi * 0.82, -math.pi * 0.82):
        p = (
            end[0] + head_len * math.cos(angle + delta),
            end[1] + head_len * math.sin(angle + delta),
        )
        draw.line([end, p], fill=fill, width=width)


def save_triangle(path: Path) -> None:
    img = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(img)
    title_f = font(46, True)
    label_f = font(32, True)
    body_f = font(24)
    small_f = font(20)

    draw.text((70, 45), "Project Management Triangle", font=title_f, fill=DARK)
    pts = [(750, 150), (265, 725), (1235, 725)]
    draw.line([pts[0], pts[1], pts[2], pts[0]], fill=BLUE_DARK, width=8)
    rounded_box(draw, (565, 110, 935, 205), BLUE_LIGHT, BLUE_DARK)
    rounded_box(draw, (85, 675, 445, 770), GREEN_LIGHT, GREEN)
    rounded_box(draw, (1055, 675, 1415, 770), AMBER_LIGHT, AMBER)
    rounded_box(draw, (560, 405, 940, 515), GRAY_LIGHT, GRAY)
    draw_wrapped_centered(draw, (565, 110, 935, 205), "Scope\nwhat must be delivered", label_f)
    draw_wrapped_centered(draw, (85, 675, 445, 770), "Time\nschedule and deadlines", label_f)
    draw_wrapped_centered(draw, (1055, 675, 1415, 770), "Cost\nbudget and resources", label_f)
    draw_wrapped_centered(draw, (560, 405, 940, 515), "Quality\nfitness for use", label_f)
    notes = [
        ("Increasing scope normally requires more time, more cost, or lower quality.", 80, 805),
        ("Compressing time raises cost or reduces what can be validated.", 80, 835),
        ("Cost limits force trade-offs in scope, pace, staffing, and assurance depth.", 80, 865),
    ]
    for txt, x, y in notes:
        draw.text((x, y), txt, font=small_f, fill=MUTED)
    draw.text((610, 550), "Quality is the outcome of managed trade-offs.", font=body_f, fill=DARK)
    img.save(path)


def save_pipeline(path: Path) -> None:
    img = Image.new("RGB", (1700, 720), "white")
    draw = ImageDraw.Draw(img)
    title_f = font(42, True)
    step_f = font(23, True)
    body_f = font(18)
    draw.text((60, 40), "Hybrid novelty and similarity detection pipeline", font=title_f, fill=DARK)
    steps = [
        ("New thesis topic", "title, abstract, keywords"),
        ("Preprocessing", "normalization and tokens"),
        ("Embedding search", "cosine similarity top-k"),
        ("Keyword analysis", "unique and common terms"),
        ("Domain rarity", "rarity by field"),
        ("Novelty score", "weighted hybrid score"),
        ("Decision label", "duplicate, near duplicate, novel"),
    ]
    x = 60
    y = 210
    w = 200
    h = 132
    gap = 34
    for i, (name, desc) in enumerate(steps):
        fill = BLUE_LIGHT if i in (2, 5, 6) else GRAY_LIGHT
        outline = BLUE_DARK if i in (2, 5, 6) else BORDER
        rounded_box(draw, (x, y, x + w, y + h), fill, outline, radius=18)
        draw_wrapped_centered(draw, (x + 12, y + 22, x + w - 12, y + 62), name, step_f)
        draw_wrapped_centered(draw, (x + 12, y + 76, x + w - 12, y + h - 14), desc, body_f, fill=MUTED)
        if i < len(steps) - 1:
            arrow(draw, (x + w + 4, y + h // 2), (x + w + gap - 8, y + h // 2), BLUE_DARK, width=4)
        x += w + gap
    formula = (
        "NoveltyScore = 0.60 * semantic distance + 0.25 * keyword uniqueness "
        "+ 0.15 * domain rarity"
    )
    draw.text((120, 505), formula, font=font(27, True), fill=DARK)
    draw.text(
        (120, 555),
        "Thresholds from the repository: similarity >= 0.90 means duplicate, >= 0.75 means near duplicate.",
        font=font(22),
        fill=MUTED,
    )
    img.save(path)


def save_wbs(path: Path) -> None:
    img = Image.new("RGB", (1800, 1180), "white")
    draw = ImageDraw.Draw(img)
    title_f = font(42, True)
    box_f = font(22, True)
    small_f = font(17)
    draw.text((70, 45), "Work Breakdown Structure for the novelty model project", font=title_f, fill=DARK)
    root = (545, 125, 1255, 210)
    rounded_box(draw, root, BLUE_LIGHT, BLUE_DARK)
    draw_wrapped_centered(draw, root, "1. University novelty screening module", font(25, True))
    groups = [
        ("1.1 Governance", ["Charter", "Stakeholder map", "Risk log"]),
        ("1.2 Data", ["Access approval", "Topic cleaning", "Labeled pairs"]),
        ("1.3 Model", ["Embeddings", "Score calibration", "Threshold evaluation"]),
        ("1.4 Integration", ["pgvector schema", "FastAPI service", "UI workflow"]),
        ("1.5 Quality", ["Unit tests", "Performance tests", "Security review"]),
        ("1.6 Rollout", ["Training", "Pilot", "Deployment"]),
    ]
    col_w = 255
    top_y = 330
    start_x = 72
    gap = 31
    for idx, (heading, children) in enumerate(groups):
        x = start_x + idx * (col_w + gap)
        rect = (x, top_y, x + col_w, top_y + 92)
        rounded_box(draw, rect, GRAY_LIGHT, BORDER, radius=14)
        draw_wrapped_centered(draw, rect, heading, box_f)
        arrow(draw, ((root[0] + root[2]) // 2, root[3]), (x + col_w // 2, top_y), BLUE_DARK, width=3)
        yy = top_y + 140
        for child in children:
            crect = (x + 15, yy, x + col_w - 15, yy + 70)
            rounded_box(draw, crect, "white", BORDER, radius=12, width=2)
            draw_wrapped_centered(draw, crect, child, small_f)
            arrow(draw, (x + col_w // 2, yy - 48), (x + col_w // 2, yy), GRAY, width=2)
            yy += 103
    draw.text(
        (70, 1060),
        "The WBS separates managerial work from data, model, platform, quality, and adoption work packages.",
        font=font(24),
        fill=MUTED,
    )
    img.save(path)


SCHEDULE = [
    ("A", "Initiation and requirements", 1, [], 0, 1, 0, 1, 0, True),
    ("B", "Data audit and access approvals", 2, ["A"], 1, 3, 1, 3, 0, True),
    ("C", "Corpus preparation and labeling", 3, ["B"], 3, 6, 5, 8, 2, False),
    ("D", "Baseline model adaptation", 2, ["B"], 3, 5, 3, 5, 0, True),
    ("E", "Hybrid score calibration", 3, ["C", "D"], 6, 9, 8, 11, 2, False),
    ("F", "Backend API and pgvector", 3, ["D"], 5, 8, 5, 8, 0, True),
    ("G", "UI workflow and dashboard", 3, ["F"], 8, 11, 8, 11, 0, True),
    ("H", "QA, security, performance tests", 2, ["E", "F", "G"], 11, 13, 11, 13, 0, True),
    ("I", "Pilot with coordinators", 2, ["H"], 13, 15, 13, 15, 0, True),
    ("J", "Deployment and training", 1, ["I"], 15, 16, 15, 16, 0, True),
    ("K", "Closure and final report", 1, ["J"], 16, 17, 16, 17, 0, True),
]


def save_gantt(path: Path) -> None:
    img = Image.new("RGB", (1800, 1180), "white")
    draw = ImageDraw.Draw(img)
    title_f = font(42, True)
    label_f = font(22, True)
    small_f = font(16)
    draw.text((60, 42), "Gantt chart and critical path", font=title_f, fill=DARK)
    draw.text((60, 96), "Critical path: A - B - D - F - G - H - I - J - K, total duration 17 weeks.", font=font(23), fill=MUTED)
    x0 = 560
    y0 = 170
    week_w = 63
    row_h = 70
    label_x = 60
    for week in range(1, 18):
        x = x0 + (week - 1) * week_w
        fill = GRAY_LIGHT if week % 2 else "white"
        draw.rectangle((x, y0, x + week_w, y0 + row_h * (len(SCHEDULE) + 1)), fill=fill, outline=BORDER)
        draw_centered(draw, (x + week_w // 2, y0 + 26), str(week), small_f, DARK)
    draw.text((label_x, y0 + 12), "Activity", font=label_f, fill=DARK)
    for i, (code, name, duration, _preds, es, ef, _ls, _lf, slack, critical) in enumerate(SCHEDULE):
        y = y0 + (i + 1) * row_h
        draw.line((60, y, x0 + week_w * 17, y), fill=BORDER, width=2)
        draw.text((label_x, y + 18), f"{code}. {name}", font=small_f, fill=DARK)
        bar_x1 = x0 + es * week_w + 8
        bar_x2 = x0 + ef * week_w - 8
        bar_y1 = y + 17
        bar_y2 = y + 48
        color = BLUE if critical else GREEN
        draw.rounded_rectangle((bar_x1, bar_y1, bar_x2, bar_y2), radius=10, fill=color)
        if slack:
            sx1 = x0 + ef * week_w + 5
            sx2 = x0 + (ef + slack) * week_w - 5
            draw.line((sx1, y + 33, sx2, y + 33), fill=AMBER, width=5)
            draw.text((sx1 + 4, y + 39), f"{slack}w slack", font=font(13), fill=AMBER)
        draw.text((bar_x1 + 8, bar_y1 + 5), f"{duration}w", font=font(14, True), fill="white")
    legend_y = 1040
    draw.rounded_rectangle((80, legend_y, 220, legend_y + 35), radius=10, fill=BLUE)
    draw.text((240, legend_y + 6), "Critical activity", font=font(20), fill=DARK)
    draw.rounded_rectangle((520, legend_y, 660, legend_y + 35), radius=10, fill=GREEN)
    draw.text((680, legend_y + 6), "Non-critical activity", font=font(20), fill=DARK)
    draw.line((980, legend_y + 18, 1120, legend_y + 18), fill=AMBER, width=5)
    draw.text((1140, legend_y + 6), "Available slack", font=font(20), fill=DARK)
    img.save(path)


def save_budget_cashflow(path: Path) -> None:
    img = Image.new("RGB", (1600, 950), "white")
    draw = ImageDraw.Draw(img)
    draw.text((60, 45), "Budget structure and planned cashflow", font=font(42, True), fill=DARK)
    categories = [
        ("Personnel", 32388, BLUE),
        ("Infrastructure", 4000, GREEN),
        ("Software and data", 800, AMBER),
        ("Training", 1000, RED),
        ("Contingency", 5728, "#7C3AED"),
    ]
    total = sum(v for _, v, _ in categories)
    x = 110
    y = 165
    max_w = 700
    for label, value, color in categories:
        w = int(max_w * value / max(v for _, v, _ in categories))
        draw.text((x, y + 7), label, font=font(22, True), fill=DARK)
        draw.rounded_rectangle((x + 220, y, x + 220 + w, y + 38), radius=9, fill=color)
        draw.text((x + 940, y + 6), f"${value:,.0f}", font=font(22, True), fill=DARK)
        y += 70
    draw.text((110, 565), f"Total planned budget: ${total:,.0f}", font=font(30, True), fill=DARK)
    cash = [7200, 10400, 10800, 7000, 2788]
    labels = ["M1", "M2", "M3", "M4", "M5"]
    base_x = 1060
    base_y = 750
    chart_w = 420
    chart_h = 470
    draw.text((980, 165), "Monthly baseline cashflow", font=font(28, True), fill=DARK)
    max_cash = max(cash)
    for i, value in enumerate(cash):
        bx = base_x + i * 78
        bh = int(chart_h * value / max_cash)
        draw.rectangle((bx, base_y - bh, bx + 46, base_y), fill=BLUE)
        draw.text((bx - 4, base_y + 12), labels[i], font=font(18, True), fill=DARK)
        draw.text((bx - 16, base_y - bh - 28), f"${value/1000:.1f}k", font=font(15), fill=MUTED)
    draw.line((base_x - 40, base_y, base_x + chart_w, base_y), fill=DARK, width=3)
    draw.line((base_x - 40, base_y - chart_h, base_x - 40, base_y), fill=DARK, width=3)
    draw.text((980, 830), "Contingency reserve is held outside monthly baseline and released by change control.", font=font(18), fill=MUTED)
    img.save(path)


def save_kpi(path: Path) -> None:
    img = Image.new("RGB", (1600, 820), "white")
    draw = ImageDraw.Draw(img)
    draw.text((60, 45), "Success KPIs for the pilot", font=font(42, True), fill=DARK)
    rows = [
        ("Duplicate detection F1", ">= 0.82", "validation set"),
        ("Near duplicate recall", ">= 0.75", "reviewer labeled pairs"),
        ("Top-k latency", "<= 1.5 sec", "API and pgvector test"),
        ("Manual review time saved", ">= 25%", "coordinator time study"),
        ("User acceptance", ">= 4.0 / 5", "pilot survey"),
        ("False high risk rate", "<= 8%", "human review audit"),
    ]
    x0, y0 = 90, 160
    col_w = [430, 250, 590]
    row_h = 78
    headers = ["Metric", "Target", "Measurement"]
    x = x0
    for i, head in enumerate(headers):
        draw.rectangle((x, y0, x + col_w[i], y0 + row_h), fill=BLUE_DARK, outline=BORDER, width=2)
        draw.text((x + 18, y0 + 24), head, font=font(24, True), fill="white")
        x += col_w[i]
    for r, row in enumerate(rows):
        y = y0 + row_h * (r + 1)
        fill = GRAY_LIGHT if r % 2 else "white"
        x = x0
        for c, value in enumerate(row):
            draw.rectangle((x, y, x + col_w[c], y + row_h), fill=fill, outline=BORDER, width=2)
            draw_wrapped_centered(draw, (x + 8, y + 8, x + col_w[c] - 8, y + row_h - 8), value, font(20, c == 1), fill=DARK)
            x += col_w[c]
    img.save(path)


def set_run_font(run, name: str = "Calibri", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = BLACK


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color.replace("#", ""))


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    width_dxa = int(sum(widths) * 1440)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths[idx] * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, header_fill: str = BLUE_DARK, header_text_white: bool = True) -> None:
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    set_run_font(run, size=9.5 if row_idx else 10, bold=row_idx == 0)
                    if row_idx == 0 and header_text_white:
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F7F9FC")


def add_table(doc: Document, rows: list[list[str]], widths: list[float], header_fill: str = BLUE_DARK):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    set_table_geometry(table, widths)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, len(row) - 1) else WD_ALIGN_PARAGRAPH.LEFT
    style_table(table, header_fill=header_fill)
    return table


def add_para(doc: Document, text: str = "", style: str | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    if style is None or style == "Normal":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, size=10.7)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, size=10.7)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(16 if level == 1 else 10)
    fmt.space_after = Pt(6)
    fmt.keep_with_next = True
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=9.5)
    run.italic = True
    run.font.color.rgb = RGBColor(75, 85, 99)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(1))
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)

    for level in (1, 2, 3):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.color.rgb = BLACK
        style.font.bold = True
        style.font.size = Pt(16 if level == 1 else 13 if level == 2 else 12)

    header = section.header.paragraphs[0]
    header.text = "Assignment 1 Project Management and Financial Justification"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=9)
        run.font.color.rgb = RGBColor(75, 85, 99)

    footer = section.footer.paragraphs[0]
    footer.text = "Hybrid novelty similarity detection model"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, size=9)
        run.font.color.rgb = RGBColor(75, 85, 99)


def generate_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    paths = {
        "triangle": ASSET_DIR / "project_triangle.png",
        "pipeline": ASSET_DIR / "model_pipeline.png",
        "wbs": ASSET_DIR / "wbs.png",
        "gantt": ASSET_DIR / "gantt_cpm.png",
        "budget": ASSET_DIR / "budget_cashflow.png",
        "kpi": ASSET_DIR / "kpi_table.png",
    }
    save_triangle(paths["triangle"])
    save_pipeline(paths["pipeline"])
    save_wbs(paths["wbs"])
    save_gantt(paths["gantt"])
    save_budget_cashflow(paths["budget"])
    save_kpi(paths["kpi"])
    return paths


def build_document() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = generate_assets()
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Project Management and Financial Justification for a Novelty Similarity Detection Model")
    set_run_font(r, size=24, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    sr = subtitle.add_run("Assignment 1 written paper and practical project case")
    set_run_font(sr, size=13)

    meta_rows = [
        ["Course topic", "Project Management and Financial Justification"],
        ["Practical case", "Hybrid thesis topic novelty and duplicate detection model"],
        ["Project source", "Repository README and modules in src/thesis_novelty"],
        ["Submission format", "Written paper, 10 to 15 pages, with supporting presentation"],
        ["Currency basis", "Illustrative 2026 USD planning budget"],
    ]
    add_table(doc, [["Field", "Value"], *meta_rows], [1.8, 4.7], header_fill="1F2937")
    add_para(
        doc,
        "This paper applies the concepts from Lecture No 2 to a practical IT and R&D project: "
        "implementation of a hybrid model that assesses thesis topic novelty and detects duplicate "
        "or near duplicate submissions. The model combines semantic similarity, keyword uniqueness, "
        "domain rarity, and threshold based classification. The work treats the model as a module "
        "inside a university thesis management platform, so the project has technical, managerial, "
        "financial, and adoption risks.",
    )
    doc.add_page_break()

    add_heading(doc, "Executive Summary", 1)
    add_para(
        doc,
        "The proposed project implements a novelty screening module that helps academic coordinators "
        "identify duplicate and near duplicate thesis topics before final approval. The model in the "
        "repository already provides a clear hybrid pipeline: preprocessing, sentence embeddings, "
        "cosine similarity search, keyword uniqueness, domain rarity, novelty scoring, and a practical "
        "classification label. The project management task is to convert that research prototype into "
        "a controlled pilot and production ready service.",
    )
    add_para(
        doc,
        "The recommended plan uses a 17 week schedule with a critical path through requirements, data "
        "approvals, baseline adaptation, API and pgvector integration, user workflow development, QA, "
        "pilot, deployment, and closure. A bottom up budget estimates a direct cost of $38,188 and a "
        "15 percent contingency reserve of $5,728, giving a total planned budget of $43,916. The "
        "financial case depends on reducing manual review time, avoiding duplicate topic remediation, "
        "and creating a reusable analytics capability for future academic quality assurance work.",
    )
    add_para(
        doc,
        "The main recommendation is to run the project as a stage gated R&D implementation rather than "
        "as a simple software installation. The model must not automatically reject a student topic. "
        "It should flag high risk similarity cases, explain the reason, and leave the academic decision "
        "to a human reviewer. That governance choice protects quality, reduces legal and academic risk, "
        "and gives the team measurable data for threshold calibration.",
    )

    add_heading(doc, "Evaluation Criteria Coverage", 2)
    coverage_rows = [
        ["Criterion", "How this paper addresses it"],
        ["Lecture content", "Summarizes project management, scheduling, budgeting, quality, tools, roles, R&D, and financial justification."],
        ["Application tools", "Uses WBS, Gantt chart, CPM, bottom up budget, cashflow, KPI targets, and a financial case."],
        ["Critical analysis", "Identifies project, technical, financial, quality, adoption, and governance risks with mitigations."],
        ["Originality", "Applies the lecture to a real hybrid novelty similarity model from the repository."],
        ["Formatting", "Includes visual diagrams, tables, structured sections, and a concise slide deck."],
    ]
    add_table(doc, coverage_rows, [1.75, 4.75], header_fill="1F2937")
    doc.add_page_break()

    add_heading(doc, "1 Analytical Review of Lecture No 2", 1)
    add_heading(doc, "1.1 Key Concepts of Project Management", 2)
    add_para(
        doc,
        "Project management is the disciplined organization of work so that a team can deliver a defined "
        "result under constraints. Lecture No 2 connects the managerial view of a project with practical "
        "control mechanisms: scope definition, scheduling, budgeting, quality management, role assignment, "
        "tool use, R&D uncertainty, and financial justification. A project differs from routine operations "
        "because it has a temporary structure, a defined objective, a set of deliverables, and a closure "
        "condition. In an IT or research setting, those deliverables include not only code or experiments "
        "but also data access, validation evidence, documentation, user training, and deployment readiness.",
    )
    add_para(
        doc,
        "The starting point is the project charter. It explains the business or research problem, defines "
        "the intended output, names stakeholders, and sets success criteria. For the novelty detection "
        "case, the charter would state that the university wants to reduce duplicate thesis topics and "
        "support a faster approval process. The charter must also define boundaries. The project can "
        "screen topics and provide evidence, but academic approval remains a human decision.",
    )
    add_para(
        doc,
        "A Work Breakdown Structure decomposes the project into manageable work packages. It does not "
        "start from the calendar. It starts from the deliverables and asks what work must exist for each "
        "deliverable to be complete. This approach helps prevent hidden work, such as data cleaning, "
        "security review, threshold calibration, and training, from appearing late in the project when "
        "they are more expensive to handle.",
    )

    add_heading(doc, "1.2 Scheduling, Budgeting, Quality, Tools, and Roles", 2)
    add_para(
        doc,
        "Scheduling translates the WBS into a sequence of activities with durations, dependencies, and "
        "milestones. A Gantt chart makes that sequence visible to the project team and sponsor. Critical "
        "Path Method analysis goes further by identifying the longest dependency chain that determines "
        "the earliest possible finish date. Activities on the critical path have zero slack, so delay in "
        "one of them delays the project unless the manager changes resources, scope, or sequence.",
    )
    add_para(
        doc,
        "Budgeting estimates the resources needed to complete each work package. The bottom up method "
        "fits implementation projects because it connects each cost to a deliverable or activity. PERT "
        "estimation is useful when uncertainty is high because it combines optimistic, most likely, and "
        "pessimistic estimates. Both methods should include a contingency reserve for known uncertainty. "
        "A separate management reserve may be held by the sponsor for unknown events, but it should not "
        "hide weak planning.",
    )
    add_para(
        doc,
        "Quality management defines what good means before work begins. In software and R&D projects, "
        "quality includes accuracy, reliability, maintainability, security, performance, reproducibility, "
        "and user acceptance. The team must agree on measurable criteria, such as F1 score, recall for "
        "near duplicates, API latency, false high risk rate, and reviewer satisfaction. Quality control "
        "checks whether outputs meet those criteria. Quality assurance checks whether the process can "
        "reliably produce them.",
    )
    add_para(
        doc,
        "Common project tools include WBS dictionaries, Gantt charts, CPM networks, risk registers, "
        "issue logs, budget trackers, Kanban or Scrum boards, test reports, dashboards, and post project "
        "lessons learned. The main roles in a project like this are sponsor, project manager, product "
        "owner, data owner, ML engineer, backend engineer, frontend or UX specialist, QA engineer, "
        "domain reviewers, and end users such as academic coordinators. R&D work adds uncertainty "
        "because model performance may depend on data quality, language coverage, labeling consistency, "
        "and threshold behavior.",
    )

    doc.add_picture(str(assets["triangle"]), width=Inches(6.45))
    add_caption(doc, "Figure 1. Project Management Triangle: scope, time, cost, and quality are interdependent constraints.")
    add_para(
        doc,
        "The Project Management Triangle explains why a project manager cannot improve every parameter "
        "at once without additional capacity. Expanding scope usually requires more time or cost. "
        "Compressing time usually requires more staffing, narrower scope, or lower assurance depth. "
        "Reducing cost limits the resources available for engineering and validation. Quality sits at "
        "the center because it is affected by all three constraints. In this case, shortening the pilot "
        "would save calendar time, but it would reduce evidence about false positives and near duplicate "
        "recall, which are quality risks.",
    )

    add_heading(doc, "1.3 Financial Justification in Project Decisions", 2)
    add_para(
        doc,
        "Financial justification connects project work to economic value. It should state the objective, "
        "the investment, the timing of cash outflows, measurable benefits, operating costs, and decision "
        "metrics such as payback period or return on investment. Financial justification should not only "
        "show that a project is affordable. It should show why the project deserves resources compared "
        "with other possible uses of the budget. For a university technology project, benefits may include "
        "administrative time saved, fewer duplicate thesis approvals, faster review cycles, and stronger "
        "quality assurance for accreditation or grant reporting.",
    )
    add_para(
        doc,
        "A good financial case also states its assumptions. In R&D projects, many benefits are estimated "
        "rather than guaranteed. The manager should therefore link funding to milestones and quality "
        "gates. If the pilot fails to reach target recall or generates too many false high risk flags, "
        "the sponsor should pause further rollout and use the remaining budget for data improvement "
        "or threshold redesign.",
    )

    add_heading(doc, "2 Practical Application", 1)
    add_heading(doc, "2.1 Project Context and Model Basis", 2)
    add_para(
        doc,
        "The practical application uses the repository novelty similarity detection model as the basis "
        "for an implementation project. The repository describes the project as a module for an "
        "intelligent platform that manages bachelor and master thesis projects. Its specific focus is "
        "novelty assessment and duplicate prevention. The implemented model is a hybrid pipeline rather "
        "than a single similarity score. It combines semantic similarity, keyword uniqueness, and domain "
        "rarity into a final novelty score.",
    )
    add_para(
        doc,
        "The model formula in the README is: NoveltyScore = alpha * (1 - max_semantic_similarity) + "
        "beta * keyword_uniqueness + gamma * domain_rarity. Default weights are alpha = 0.60, beta = "
        "0.25, and gamma = 0.15. The classification thresholds are similarity >= 0.90 for duplicate, "
        "similarity >= 0.75 for near duplicate, and otherwise novel. The repository also includes a "
        "FastAPI service, CLI inference, sample data, threshold evaluation script, tests, and a pgvector "
        "schema for a production style vector database.",
    )
    doc.add_picture(str(assets["pipeline"]), width=Inches(6.45))
    add_caption(doc, "Figure 2. Practical model basis from the repository: a hybrid explainable novelty pipeline.")

    add_heading(doc, "2.2 Project Objectives and Scope", 2)
    add_para(
        doc,
        "The project objective is to implement and validate the novelty screening module in a university "
        "thesis management workflow. The immediate product is not a fully autonomous decision system. "
        "It is a decision support service that helps coordinators and supervisors see whether a proposed "
        "topic is semantically close to existing topics, which keywords are rare or common, and whether "
        "the system classifies the case as duplicate, near duplicate, or novel.",
    )
    add_para(doc, "The in scope deliverables are:")
    for item in [
        "Cleaned corpus of historical thesis topics with title, abstract, keywords, domain, supervisor, year, and language fields.",
        "Validated model configuration with documented weights, thresholds, and evaluation results.",
        "FastAPI service connected to a vector search backend for top-k similarity search.",
        "Coordinator facing workflow that displays similar topics, novelty components, and a reviewer decision field.",
        "Quality dashboard with model metrics, latency, false high risk rate, and reviewer feedback.",
        "Pilot report, deployment checklist, training material, and final lessons learned report.",
    ]:
        add_bullet(doc, item)
    add_para(
        doc,
        "Out of scope are automatic rejection of student topics, supervisor assignment optimization, "
        "full thesis lifecycle management, and replacement of academic committee review. Keeping those "
        "items out of scope protects schedule and quality because the model needs a narrow first use case "
        "before wider platform functions are added.",
    )
    doc.add_page_break()

    add_heading(doc, "2.3 Work Breakdown Structure", 2)
    doc.add_picture(str(assets["wbs"]), width=Inches(6.45))
    add_caption(doc, "Figure 3. WBS for implementing the novelty screening module.")
    wbs_rows = [
        ["WBS code", "Work package", "Main deliverables"],
        ["1.1", "Governance and project control", "Charter, stakeholder map, schedule baseline, risk register, change control process"],
        ["1.2", "Data and compliance", "Data access approval, cleaned historical topics, labeled validation pairs, privacy notes"],
        ["1.3", "Model development", "Embedding backend, score calibration, threshold evaluation, reproducible experiment report"],
        ["1.4", "Platform integration", "pgvector schema, API endpoint, UI workflow, logging and monitoring"],
        ["1.5", "Quality and validation", "Unit tests, integration tests, performance tests, security review, acceptance criteria"],
        ["1.6", "Pilot and rollout", "Training, pilot feedback, production deployment, closure report"],
    ]
    add_table(doc, wbs_rows, [1.05, 2.0, 3.45], header_fill="1F2937")
    add_para(
        doc,
        "The WBS is built around deliverables, not around people. This matters because one person may "
        "support several work packages, and one work package may need several roles. For example, model "
        "development requires the ML engineer, domain reviewers, and the project manager because threshold "
        "choice is both a technical and governance decision.",
    )

    add_heading(doc, "2.4 Gantt Chart and Critical Path", 2)
    add_para(
        doc,
        "The project is planned as a 17 week implementation. The schedule starts with requirements and "
        "data approvals because a model project cannot proceed reliably without access to historical topic "
        "data. Model adaptation and platform integration then run in a controlled sequence. The pilot starts "
        "only after QA, security, and performance tests confirm that the service can support real review work.",
    )
    doc.add_picture(str(assets["gantt"]), width=Inches(6.45))
    add_caption(doc, "Figure 4. Gantt chart and CPM view. Blue bars are critical path activities.")
    cpm_rows = [["ID", "Activity", "Dur.", "Pred.", "ES", "EF", "LS", "LF", "Slack", "Critical"]]
    for code, name, duration, preds, es, ef, ls, lf, slack, critical in SCHEDULE:
        cpm_rows.append([
            code,
            name,
            str(duration),
            ",".join(preds) if preds else "-",
            str(es),
            str(ef),
            str(ls),
            str(lf),
            str(slack),
            "Yes" if critical else "No",
        ])
    add_table(doc, cpm_rows, [0.45, 2.55, 0.45, 0.55, 0.35, 0.35, 0.35, 0.35, 0.55, 0.45], header_fill="1F2937")
    add_para(
        doc,
        "The critical path is A - B - D - F - G - H - I - J - K. It determines the 17 week finish date. "
        "Corpus preparation and hybrid calibration have two weeks of slack because the integration stream "
        "through API, pgvector, and UI workflow is longer. This slack should not be wasted. It should be used "
        "to improve labeled pairs and check whether thresholds perform consistently across languages and domains.",
    )
    doc.add_page_break()

    add_heading(doc, "2.5 Bottom Up Budget with Contingency Reserve", 2)
    add_para(
        doc,
        "The budget uses a bottom up estimate because the work packages are specific enough to estimate "
        "from role effort and direct costs. The figures are illustrative 2026 USD planning values. They "
        "should be replaced by institutional salary rates and procurement rules before formal approval.",
    )
    budget_rows = [
        ["Cost item", "Basis", "Amount"],
        ["Project manager / analyst", "0.25 FTE x 4.25 months x $4,000", "$4,250"],
        ["Data engineer", "0.50 FTE x 3.00 months x $3,200", "$4,800"],
        ["ML engineer", "0.75 FTE x 3.50 months x $4,500", "$11,813"],
        ["Backend engineer", "0.50 FTE x 3.00 months x $4,000", "$6,000"],
        ["Frontend and UX", "0.30 FTE x 2.50 months x $3,500", "$2,625"],
        ["QA engineer", "0.25 FTE x 2.00 months x $2,800", "$1,400"],
        ["Domain reviewers", "60 hours x $25", "$1,500"],
        ["Cloud experiments", "GPU/CPU experiments and model evaluation", "$1,800"],
        ["Staging database and pgvector hosting", "5 months x $300", "$1,500"],
        ["Monitoring, backup, and security tools", "Fixed implementation allowance", "$700"],
        ["Annotation and workflow tools", "5 seats x 3 months x $30", "$450"],
        ["Reporting, CI, documentation tools", "Fixed allowance", "$350"],
        ["Training workshops", "Two coordinator sessions", "$600"],
        ["Training materials", "Guides and recorded walkthrough", "$400"],
        ["Direct cost subtotal", "", "$38,188"],
        ["Contingency reserve", "15 percent of direct cost", "$5,728"],
        ["Total planned budget", "", "$43,916"],
    ]
    add_table(doc, budget_rows, [2.25, 3.35, 0.9], header_fill="1F2937")
    doc.add_picture(str(assets["budget"]), width=Inches(6.45))
    add_caption(doc, "Figure 5. Budget structure and planned monthly cashflow.")
    add_para(
        doc,
        "The contingency reserve is not a hidden spending line. It is a control mechanism for known "
        "uncertainty, especially data cleaning effort, threshold recalibration, and integration defects. "
        "Reserve release should require a change request that states the cause, the affected deliverable, "
        "and the effect on scope, time, cost, or quality.",
    )

    add_heading(doc, "2.6 Cashflow, KPIs, and Financial Justification", 2)
    cash_rows = [
        ["Month", "Main work focus", "Baseline cashflow", "Cumulative"],
        ["1", "Initiation, requirements, data access", "$7,200", "$7,200"],
        ["2", "Corpus preparation, model adaptation", "$10,400", "$17,600"],
        ["3", "API, pgvector, UI workflow, calibration", "$10,800", "$28,400"],
        ["4", "QA, security, performance, pilot", "$7,000", "$35,400"],
        ["5", "Deployment, training, closure", "$2,788", "$38,188"],
    ]
    add_table(doc, cash_rows, [0.65, 3.55, 1.15, 1.15], header_fill="1F2937")
    add_para(
        doc,
        "The quantified benefit case uses conservative planning assumptions for a medium sized university. "
        "If 2,500 thesis topics are screened each year and the module saves 15 minutes of coordinator or "
        "supervisor review per topic, it saves 625 hours. At a loaded labor rate of $25 per hour, that "
        "equals $15,625 per year. If better screening prevents 60 late duplicate remediation cases per "
        "year and each case costs about $350 in administrative and supervisory time, the avoided cost is "
        "$21,000. A reusable evaluation and reporting module adds an estimated $8,000 in internal analytics "
        "value. Total annual quantified benefit is therefore $44,625.",
    )
    add_para(
        doc,
        "Annual operating support is estimated at $8,000 for hosting, monitoring, small fixes, and periodic "
        "threshold review. Net annual benefit is $36,625. Against a one time planned budget of $43,916, the "
        "simple payback period is about 1.20 years. The two year ROI is approximately 67 percent: "
        "(2 x $36,625 - $43,916) / $43,916. These figures should be validated during the pilot through "
        "time tracking, duplicate case counts, and user feedback.",
    )
    doc.add_picture(str(assets["kpi"]), width=Inches(6.45))
    add_caption(doc, "Figure 6. Pilot KPI targets for quality, usefulness, and operational performance.")
    doc.add_page_break()

    add_heading(doc, "3 Critical Analysis", 1)
    add_heading(doc, "3.1 Main Risks and Challenges", 2)
    add_para(
        doc,
        "The main project risk is that similarity detection looks simple to stakeholders but is difficult "
        "to govern. A high similarity score may indicate plagiarism, reuse of a common topic, normal "
        "terminology within a domain, or a legitimate continuation of earlier research. If the tool is "
        "treated as an automatic judgment system, it may damage student trust and academic fairness. "
        "The correct managerial approach is to define it as decision support and require human review "
        "for all high risk cases.",
    )
    risk_rows = [
        ["Risk", "Impact", "Mitigation"],
        ["Historical data quality is inconsistent", "Model returns noisy similarity results and weak keyword statistics", "Data audit, cleaning rules, minimum fields, and periodic corpus quality checks"],
        ["False positives block novel topics", "Students and supervisors lose trust in the platform", "Human in the loop review, appeal path, threshold tuning, and false high risk KPI"],
        ["Language imbalance affects results", "Kazakh, Russian, and English topics may not perform equally", "Language tagged validation sets and multilingual embedding tests"],
        ["Data privacy constraints delay work", "Schedule slips before model development starts", "Early data owner approval, anonymization, and documented access controls"],
        ["Integration defects affect review workflow", "Coordinators avoid the tool or duplicate manual work", "API contract tests, UI pilot, logging, and staged rollout"],
        ["Budget pressure removes validation work", "Short term savings create quality and adoption failures", "Protected QA budget, reserve release rules, and sponsor sign off for scope changes"],
        ["R&D performance misses target", "Pilot cannot justify production rollout", "Stage gate after evaluation, baseline comparison, and fallback to narrowed scope"],
    ]
    add_table(doc, risk_rows, [1.75, 2.1, 2.65], header_fill="1F2937")
    add_para(
        doc,
        "These risks show why the lecture topics belong together. Scheduling without risk management "
        "would hide data approval and validation uncertainty. Budgeting without quality metrics would "
        "make the project appear cheaper than it is. Quality management without financial justification "
        "would not explain why the university should fund the pilot. The manager needs all three views.",
    )

    add_heading(doc, "3.2 Control Tools and Quality Metrics", 2)
    add_para(doc, "The recommended control system combines project controls with model controls:")
    for item in [
        "Schedule control through a weekly Gantt review, critical path status, and dependency review for data access, API integration, and pilot readiness.",
        "Budget control through monthly planned versus actual cost, reserve drawdown tracking, and approval for any scope change that affects total cost.",
        "Risk control through a live risk register with probability, impact, owner, mitigation status, and trigger date.",
        "Quality control through model metrics, API latency tests, security checks, reviewer audit, and user acceptance surveys.",
        "Governance control through a rule that the model flags risk but does not make final academic decisions.",
    ]:
        add_bullet(doc, item)
    add_para(
        doc,
        "The most important quality metrics are duplicate detection F1, near duplicate recall, false high "
        "risk rate, top-k latency, explanation completeness, and reviewer agreement. Model accuracy alone "
        "is not enough. The system must also produce an explanation that reviewers can understand, such as "
        "which historical topics were similar, what the semantic similarity was, which keywords are common, "
        "and why the domain is rare or common.",
    )
    doc.add_page_break()

    add_heading(doc, "4 Conclusions and Recommendations", 1)
    add_para(
        doc,
        "The most relevant lesson from Lecture No 2 is that project success depends on managing trade-offs "
        "explicitly. In the novelty similarity detection case, the team cannot maximize scope, speed, low "
        "cost, and quality at the same time. A short implementation without data governance and validation "
        "would be risky even if the code runs. A wider scope that also includes supervisor assignment and "
        "full thesis management would exceed the schedule and budget of the pilot. The recommended approach "
        "is therefore a narrow, measurable, stage gated implementation.",
    )
    add_para(
        doc,
        "Scheduling methods apply directly to real grant funded, corporate, and startup projects. In a grant "
        "project, the Gantt chart and CPM show whether milestones can be met before reporting deadlines. In a "
        "corporate project, the schedule helps coordinate data owners, engineering, security, and users. In a "
        "startup project, it helps founders decide which features must wait until product market evidence is "
        "available. In each case, the critical path prevents management attention from spreading evenly across "
        "all activities when only some activities determine the finish date.",
    )
    add_para(
        doc,
        "Financial justification methods also transfer across settings. A grant proposal needs a budget that "
        "connects resources to work packages and explains why the reserve is reasonable. A corporate project "
        "needs payback, ROI, operating cost, and KPI commitments. A startup project needs a burn plan and a "
        "clear reason why the work will create learning or revenue. The same discipline applies to this model: "
        "the project should be funded only if the pilot can measure review time saved, duplicate cases avoided, "
        "quality performance, and user adoption.",
    )
    add_para(
        doc,
        "The final recommendation is to approve a controlled pilot with the $43,916 planned budget, a 17 week "
        "schedule baseline, and the KPI targets shown above. Production rollout should be conditional on "
        "meeting quality gates, especially near duplicate recall, false high risk rate, and user acceptance. "
        "If those gates are not met, the team should use the remaining reserve for data improvement or "
        "threshold redesign rather than forcing deployment.",
    )

    add_heading(doc, "References and Source Basis", 1)
    for item in [
        "Lecture No 2 assignment topics: project management, scheduling, budgeting, quality management, tools, roles, R&D, and financial justification.",
        "Repository README: Hybrid Thesis Topic Novelty Assessment Model, model formula, default weights, thresholds, CLI, API, experiment, and pgvector notes.",
        "Repository source files: src/thesis_novelty/model.py, config.py, schemas.py, api.py, sample_topics.csv, sample_pairs.csv, and tests/test_model.py.",
        "Budget, cashflow, benefits, and KPI values are planning assumptions for this assignment and should be replaced with institutional data for approval.",
    ]:
        add_bullet(doc, item)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
