from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "Assignment_2_Methodology_and_Architecture_Novelty_Model.docx"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F4F7FA"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_col_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_run(p, text, bold=False, size=None, color=None, italic=False):
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Aptos"
    return run


def set_paragraph(p, before=0, after=6, line=1.08):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def remove_title_rule(p):
    """Override Word's built-in Title bottom border for a plain cover page."""
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_body(doc, text, before=0, after=6):
    p = doc.add_paragraph()
    set_paragraph(p, before, after)
    add_run(p, text, size=10.5)
    return p


def add_bold_lead(doc, lead, text):
    p = doc.add_paragraph()
    set_paragraph(p, 0, 6)
    add_run(p, lead, bold=True, size=10.5)
    add_run(p, text, size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_paragraph(p, 10 if level == 1 else 6, 4)
    add_run(p, text, bold=True, size=14 if level == 1 else 11.5, color="000000")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(p, 2, 7)
    add_run(p, text, italic=True, size=9, color="555555")
    return p


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for j, header in enumerate(headers):
        set_col_width(hdr[j], widths[j])
        set_cell_shading(hdr[j], NAVY)
        set_cell_border(hdr[j])
        set_cell_margins(hdr[j])
        hdr[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(p, 0, 0, 1.0)
        add_run(p, header, bold=True, size=font_size, color="FFFFFF")
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_col_width(cells[j], widths[j])
            set_cell_border(cells[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 1:
                set_cell_shading(cells[j], PALE)
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 or (len(row) > 2 and j == len(row)-1) else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph(p, 0, 0, 1.0)
            add_run(p, str(value), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_architecture_diagram(doc):
    rows = [
        ["Platform or reviewer UI", "FastAPI adapter", "Assessment application service", "Domain scorer and policy"],
        ["Integration API", "Authentication and audit adapter", "Embedding provider port", "Topic repository port"],
        ["Reviewer decision", "Metrics and logs", "Multilingual E5 model", "PostgreSQL plus pgvector"],
    ]
    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    colors = [["DCE6F1", "DCE6F1", "DCE6F1", "DCE6F1"], ["F4F7FA", "F4F7FA", "F4F7FA", "F4F7FA"], ["EAF2E5", "EAF2E5", "EAF2E5", "EAF2E5"]]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            set_col_width(cell, 1.58)
            set_cell_border(cell, color="9FBAD0", size="8")
            set_cell_shading(cell, colors[i][j])
            set_cell_margins(cell, top=150, bottom=150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph(p, 0, 0, 1.0)
            add_run(p, value, bold=i == 0, size=8.8)
    add_caption(doc, "Figure 1. Target modular monolith and its replaceable adapters")


def add_delivery_flow(doc):
    labels = ["0  Scope and corpus", "1  Baseline", "2  API and review", "3  Pilot", "Gate  Quality release"]
    table = doc.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for j, label in enumerate(labels):
        cell = table.cell(0, j)
        set_col_width(cell, 1.26)
        set_cell_border(cell, color="9FBAD0", size="8")
        set_cell_shading(cell, BLUE if j < 4 else "EAF2E5")
        set_cell_margins(cell, top=150, bottom=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(p, 0, 0, 1.0)
        add_run(p, label, bold=True, size=8.5)
    add_caption(doc, "Figure 2. Four short delivery increments with explicit evidence gates")


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(p, 0, 0, 1.0)
    add_run(p, "Assignment 2  Methodology and Architecture Assessment", size=8, color="666666")
    p.add_run("   |   ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def page_break(doc):
    doc.add_page_break()


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.66)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    add_footer(sec)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for style in (styles["Heading 1"], styles["Heading 2"], styles["Title"]):
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)

    # Page 1: title page
    doc.add_paragraph().paragraph_format.space_after = Pt(52)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_title_rule(p)
    set_paragraph(p, 0, 10, 1.0)
    add_run(p, "Methodology and Architecture Assessment", bold=True, size=24, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(p, 0, 24)
    add_run(p, "Hybrid Thesis Topic Novelty Assessment Model", size=15, color="000000")
    add_body(doc, "Assignment 2  Project methodology and architectural approach", after=15).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(doc, "Alisher Tokesh  CSE 2507M", after=8).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(doc, "Project basis: novelty-similarity-detection-model repository", after=8).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(190)
    add_body(doc, "October 2026", after=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_break(doc)

    # Page 2: table of contents and executive note
    add_heading(doc, "Table of Contents")
    toc = [
        ("1. Project Analysis", "3"), ("1.1 Goals and development context", "3"),
        ("1.2 Requirements and priorities", "4"), ("1.3 Constraints and risks", "4"),
        ("2. Development Methodology Assessment", "5"), ("2.1 Recommended hybrid approach", "5"),
        ("2.2 Fit, risks and alternatives", "6"), ("3. Architectural Approach Assessment", "7"),
        ("3.1 Current state and target design", "7"), ("3.2 Quality attributes and adjustments", "8"),
        ("4. Conclusions and Recommendations", "9"), ("References", "9"),
    ]
    for label, number in toc:
        p = doc.add_paragraph()
        set_paragraph(p, 0, 3)
        add_run(p, label, size=10.5)
        p.add_run(" " * max(2, 66 - len(label)))
        add_run(p, number, size=10.5)
    add_heading(doc, "Executive Position", 1)
    add_body(doc, "This report evaluates the existing Hybrid Thesis Topic Novelty Assessment Model as a practical module for an intelligent thesis-management platform. The repository already provides an explainable hybrid score, interchangeable embedding backends, a command-line and FastAPI interface, automated tests, and a PostgreSQL plus pgvector schema. The principal recommendation is a modular monolith delivered through four evidence-based increments. This approach preserves flexibility for threshold calibration and multilingual data while keeping governance, auditability, and operational cost proportionate to the approved pilot case.")
    add_bold_lead(doc, "Planning baseline from Assignment 1. ", "The methodology and architecture are assessed against the 17-week stage-gated R and D implementation plan, a direct-cost estimate of $38,188, and a 15 percent contingency reserve of $5,728, for a total planned budget of $43,916. The model is decision support: it flags high-risk similarity cases, while academic approval remains a human decision.")
    page_break(doc)

    # Page 3
    add_heading(doc, "1. Project Analysis")
    add_heading(doc, "1.1 Goals and Development Context", 2)
    add_body(doc, "The product objective is to help a university identify duplicate or near-duplicate thesis topics before approval and to provide a transparent basis for reviewer judgement. A submitted title, abstract, keywords and domain are normalized, embedded, compared with an existing corpus, and evaluated for keyword uniqueness and domain rarity. The result is a novelty score, duplicate probability, label, explanation components, and the most similar historical topics.")
    add_bold_lead(doc, "Primary goals. ", "Reduce duplicate-topic risk; support multilingual academic topics; make a recommendation explainable rather than purely opaque; preserve a reproducible baseline for research; and expose the assessment through an API that can later join a wider thesis-management platform.")
    add_bold_lead(doc, "Current implementation evidence. ", "The model calculates semantic distance, keyword uniqueness, and domain rarity with configurable weights (0.60, 0.25, 0.15). Default similarity thresholds classify duplicate at 0.90 or above and near-duplicate at 0.75 or above. The repository offers a deterministic hashing fallback for tests and an optional multilingual SentenceTransformer backend for production-like evaluation.")
    add_heading(doc, "Development Context", 2)
    add_table(doc, ["Factor", "Assessment", "Implication"], [
        ["Project type", "Applied ML and API module", "Both model quality and software quality must be demonstrated."],
        ["Users", "Academic administrators and supervisors", "Explainability and reviewer control are essential."],
        ["Data", "Historical topic corpus in several languages", "Data quality and labels are major delivery risks."],
        ["Lifecycle", "Pilot, then platform integration", "Avoid irreversible coupling to a UI or one model vendor."],
    ], [1.18, 2.0, 3.62])
    add_caption(doc, "Table 1. Context derived from the repository and the stated thesis-platform purpose")
    page_break(doc)

    # Page 4
    add_heading(doc, "1.2 Requirements and Priorities", 2)
    add_body(doc, "Priorities use MoSCoW: Must means required for a trustworthy pilot; Should means high value after the vertical slice works; Could is deliberately deferred. Quantitative non-functional targets are pilot acceptance targets proposed for planning.")
    add_table(doc, ["Type", "Priority", "Requirement", "Acceptance signal"], [
        ["Functional", "Must", "Assess a topic against a corpus and return top similar topics, score, probability and label.", "Deterministic API response with top-k evidence."],
        ["Functional", "Must", "Ingest existing topics and keep corpus embeddings consistent with the chosen model.", "Rebuild process records corpus and model version."],
        ["Functional", "Should", "Allow a reviewer to confirm or override the suggested label with a reason.", "Decision and reason are auditable."],
        ["Functional", "Should", "Evaluate thresholds on labelled duplicate, near-duplicate and novel pairs.", "Precision, recall and F1 report."],
        ["Functional", "Could", "Provide a separate administrative dashboard and retraining automation.", "Deferred beyond pilot unless demand is proven."],
        ["Nonfunctional", "Must", "Explainability and reproducibility of every assessment.", "Store score components, thresholds and model ID."],
        ["Nonfunctional", "Must", "Privacy and access control for topics and reviewer decisions.", "RBAC, encrypted transport, restricted logs."],
        ["Nonfunctional", "Should", "p95 assessment response under 1.5 seconds for a 10,000-topic pilot corpus.", "Load test using production-like embeddings."],
    ], [0.86, 0.72, 3.2, 2.02], font_size=7.55)
    add_caption(doc, "Table 2. Prioritized requirements")
    add_heading(doc, "1.3 Constraints and Risks", 2)
    add_table(doc, ["Constraint or risk", "Why it matters", "Control"], [
        ["17-week schedule and $43,916 cap", "There is limited room for rework or infrastructure complexity.", "Critical-path control, vertical slices and protected contingency."],
        ["Sparse or inconsistent historical records", "Similarity can look precise while being wrong.", "Data profiling, sampling and a labelled evaluation set."],
        ["Multilingual semantic variation", "A fallback vectorizer is insufficient for final claims.", "Evaluate multilingual E5 against a frozen benchmark."],
        ["False duplicate decision", "Can unfairly block a student's topic.", "Human-in-the-loop decision, threshold review and audit trail."],
    ], [1.72, 2.6, 2.48], font_size=8.0)
    page_break(doc)

    # Page 5
    add_heading(doc, "2. Development Methodology Assessment")
    add_heading(doc, "2.1 Recommended Hybrid Approach", 2)
    add_body(doc, "The recommended methodology is iterative incremental delivery within the 17-week stage-gated R and D plan: Scrum-style two-week planning and review cycles for product and model work, Kanban limits for data remediation and review tasks, and evidence gates at the key critical-path transitions. The gates are not a second management system. They are evidence checks because a classifier can be technically deployable while still failing its educational or fairness purpose.")
    add_delivery_flow(doc)
    add_bold_lead(doc, "Increment 0  weeks 1 to 3. ", "Complete initiation, requirements, data-access approval, privacy classification, duplicate definitions, labelled-pair protocol, baseline metrics, and a definition of done. Gate 1 requires approved data access and traceable corpus provenance.")
    add_bold_lead(doc, "Increment 1  weeks 4 to 6. ", "Deliver corpus preparation, a hashing baseline, model adaptation, explainable result contract, and test fixtures. The critical path passes from data approval through baseline adaptation into backend integration.")
    add_bold_lead(doc, "Increment 2  weeks 7 to 11. ", "Add multilingual E5 embeddings, pgvector retrieval, threshold evaluation, API integration, coordinator workflow, reviewer override, and a simple audit event. Gate 2 checks F1 by class, error samples, and latency before QA begins.")
    add_bold_lead(doc, "Increment 3  weeks 12 to 17. ", "Complete QA, security and performance testing; run the coordinator pilot; collect disagreement reasons; then deploy, train users, and close the project. Gate 3 requires documented residual risks and user acceptance before rollout.")
    add_heading(doc, "Roles and Cadence", 2)
    add_table(doc, ["Role", "Responsibility", "Cadence"], [
        ["Project manager and academic lead", "Own schedule, budget, policy, examples and reviewer workflow.", "Biweekly review"],
        ["Data engineer", "Corpus quality, access controls and labelled validation pairs.", "Weekly data review"],
        ["ML engineer", "Embeddings, benchmark, threshold calibration.", "Daily work and experiment review"],
        ["Backend engineer", "API, persistence, security and observability.", "Daily work and integration review"],
        ["Domain reviewers and QA", "Labels, challenge cases and acceptance testing.", "Weekly sample review"],
    ], [1.45, 3.4, 1.95], font_size=8.3)
    page_break(doc)

    # Page 6
    add_heading(doc, "2.2 Fit, Risks and Alternatives", 2)
    add_body(doc, "This hybrid approach fits the project because the algorithm's thresholds, corpus quality and acceptable trade-offs cannot be reliably specified at the start. Small increments make technical and policy assumptions visible early, while the Assignment 1 critical path keeps work focused on data approval, baseline adaptation, API and pgvector integration, workflow, QA, pilot, deployment and closure. A full Scrum implementation would be excessive for a cross-functional part-time team, so the recommended version retains only useful ceremonies: a short planning session, visible work board, weekly quality review, and demonstrable increment every two weeks.")
    add_table(doc, ["Criterion", "Hybrid incremental", "Waterfall", "Kanban only"], [
        ["Model and data uncertainty", "5  learn and recalibrate", "2  late feedback", "4  flexible, weaker checkpoints"],
        ["Fixed pilot deadline", "4  time-boxed slices", "4  predictable plan", "3  flow can obscure deadline"],
        ["Small team overhead", "4  lightweight roles", "4  low ceremony", "5  very low ceremony"],
        ["Governance and audit", "5  gates plus evidence", "5  document-oriented", "2  needs added controls"],
        ["Adaptability after pilot", "5  backlog and adapters", "2  change requests", "4  continuous flow"],
    ], [2.1, 1.53, 1.37, 1.8], font_size=8.15)
    add_caption(doc, "Table 3. Relative fit scored 1 low to 5 high")
    add_heading(doc, "Methodology Risks and Responses", 2)
    add_table(doc, ["Risk", "Mitigation", "Alternative trigger"], [
        ["Iterations turn into unmeasured experimentation", "Freeze benchmark; log dataset, model and threshold version; use gate criteria.", "Use a short research protocol if experiments exceed one sprint."],
        ["Stakeholders cannot attend reviews", "Asynchronous demo, decision log, named delegate.", "Adopt a fixed weekly approval window."],
        ["Kanban data tasks grow without bound", "WIP limit and explicit backlog triage.", "Split data remediation into a separately funded stream."],
        ["Institution requires fixed compliance sign-off", "Move threat model and data approval into Increment 0.", "Use a lightweight Waterfall phase only for compliance artifacts."],
    ], [1.65, 3.16, 2.0], font_size=7.9)
    add_bold_lead(doc, "Decision. ", "A pure Waterfall plan is acceptable only where the data, acceptance rules and legal controls are already stable. A pure Kanban system is preferable later for operations and corpus maintenance, but it is not sufficient alone for the first evidence-based release.")
    page_break(doc)

    # Page 7
    add_heading(doc, "3. Architectural Approach Assessment")
    add_heading(doc, "3.1 Current State and Target Design", 2)
    add_body(doc, "The implemented code is already close to a modular monolith: preprocessing, embeddings, keyword analysis, schemas, scoring model, CSV I/O, CLI and FastAPI are separate modules. The EmbeddingBackend protocol lets the system switch between hashing and SentenceTransformer encoders. The production persistence path is not yet wired into the model: the running prototype stores topics and vectors in process memory, although the repository includes a pgvector schema with a cosine HNSW index.")
    add_bold_lead(doc, "Recommended style. ", "Keep a deployable modular monolith, but formalize it as a hexagonal design. The assessment application service owns the use case; domain scoring remains independent of infrastructure; adapters implement the API, embedding provider, repository, audit logging and observability. This gives the pilot clear boundaries without the operational burden of microservices.")
    add_architecture_diagram(doc)
    add_heading(doc, "Why This Architecture Fits", 2)
    add_table(doc, ["Need", "Alignment", "Limit or decision"], [
        ["Top-k similarity at scale", "pgvector adapter and existing HNSW schema support indexed cosine search.", "Validate recall and latency before selecting index settings."],
        ["Explainable result", "Scorer returns components, similar topics and label separately from transport.", "Persist configuration and reviewer decision too."],
        ["Model replacement", "EmbeddingBackend protocol isolates encoder choice.", "Store embedding dimension and model version; re-index on change."],
        ["Small-team delivery", "One deployable service simplifies releases and monitoring.", "Do not split services until load or team boundaries justify it."],
    ], [1.5, 3.35, 1.96], font_size=8.0)
    page_break(doc)

    # Page 8
    add_heading(doc, "3.2 Quality Attributes and Adjustments", 2)
    add_table(doc, ["Attribute", "Current support", "Required adjustment"], [
        ["Reliability", "Unit tests cover duplicate and unrelated cases; API has health endpoint.", "Add integration tests for storage, restart recovery, timeouts and malformed corpus data."],
        ["Security and privacy", "No authentication or audit layer is visible in current API.", "Add platform RBAC, TLS, least-privilege database role, redacted logs and retention policy."],
        ["Performance", "In-memory ranking is simple; schema contains HNSW index.", "Use pgvector retrieval, batch embedding, bounded top-k and measured load tests."],
        ["Maintainability", "Separate modules and typed dataclasses aid change.", "Introduce repository and audit ports; document ADRs and contract versioning."],
        ["Reproducibility", "Configurable weights and deterministic fallback exist.", "Version datasets, embeddings, thresholds and evaluation reports together."],
    ], [1.15, 2.85, 2.81], font_size=7.9)
    add_caption(doc, "Table 4. Quality attribute assessment")
    add_heading(doc, "Justified Architectural Adjustments", 2)
    add_bold_lead(doc, "1. Persistence boundary. ", "Replace direct list ownership in the service deployment with a TopicRepository port and a pgvector implementation. It should return candidate records and vector-search metadata; scoring remains in the application/domain layer.")
    add_bold_lead(doc, "2. Decision audit. ", "Persist assessment ID, corpus version, embedding model, threshold and weight configuration, components, returned candidates, reviewer decision and reason. This enables appeals, calibration analysis and reproducible research claims.")
    add_bold_lead(doc, "3. Controlled asynchronous work. ", "Ingestion, embedding rebuilds and evaluations should run as jobs rather than inside a request. The assessment path stays synchronous and bounded; rebuild failures cannot interrupt normal review.")
    add_bold_lead(doc, "4. Observability and policy. ", "Record latency, error rate, label distribution, reviewer disagreement and drift in similarity distributions. Logs must avoid storing raw abstracts unless the approved data policy permits it.")
    add_bold_lead(doc, "5. Avoid premature microservices. ", "Split the embedding or evaluation worker later only when independent scaling, separate deployment cadence, or security isolation produces a measured benefit. For the pilot, a single service with explicit ports is the lower-risk choice.")
    page_break(doc)

    # Page 9
    add_heading(doc, "4. Conclusions and Recommendations")
    add_body(doc, "The recommended methodology and architecture are well aligned with the project goals. The hybrid score and multilingual retrieval problem require learning cycles, but the effect on students and academic approval means those cycles need visible evidence, stable evaluation data and human review. A compact iterative process with quality gates therefore offers a better fit than a rigid linear plan or unstructured continuous flow.")
    add_body(doc, "A modular monolith with hexagonal boundaries also matches the present codebase and the likely pilot constraints. It uses existing seams around embeddings and API delivery, adopts the repository's pgvector direction for scalable search, and reserves microservice extraction for a demonstrated need. The strongest missing capabilities are not another scoring formula; they are persistent corpus handling, policy-controlled auditability, security controls, benchmark governance and production observability.")
    add_heading(doc, "Prioritized Recommendations", 2)
    add_table(doc, ["When", "Recommendation", "Expected result"], [
        ["Weeks 1 to 3", "Approve data access, define duplicate policy, freeze labelled benchmark and metric thresholds.", "A governable baseline and an executable critical-path start."],
        ["Weeks 4 to 6", "Prepare the corpus, adapt the baseline model and document threshold evidence.", "Reproducible model behaviour before integration."],
        ["Weeks 7 to 11", "Deliver pgvector adapter, API, coordinator workflow, audit event and measurable retrieval.", "Production-like assessment on the integration critical path."],
        ["Weeks 12 to 17", "Complete QA and security review, pilot, deployment, training and closure within the $43,916 baseline.", "Evidence for scale-up or targeted revision."],
    ], [1.05, 3.85, 1.91], font_size=8.1)
    add_heading(doc, "References", 2)
    refs = [
        "1. Hybrid Thesis Topic Novelty Assessment Model. README.md. Repository documentation, accessed October 2026.",
        "2. thesis_novelty.model, config, embeddings, api and schemas modules. Source code reviewed for this report.",
        "3. sql/pgvector_schema.sql. Proposed PostgreSQL and pgvector persistence schema reviewed for this report.",
        "4. Assignment 2 brief supplied by the course. Planning assumptions in this report were created solely for the academic case study.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        set_paragraph(p, 0, 3, 1.0)
        add_run(p, ref, size=8.8)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
